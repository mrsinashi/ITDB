from io import BytesIO
import re

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func

from db import SessionLocal
from models import (
    Choice,
    Computer,
    ComputerPerson,
    Location,
    Person,
    VacuumAccount,
    VacuumAccountComputer,
    VacuumAccountPerson,
)

from api_import import (
    HEADER_MAPPING,
    IMPORTANT_FIELDS,
    cell_to_text,
    clean_text,
    detect_header,
    normalize_header,
    parse_drive_text,
    parse_int,
    parse_ips,
    parse_macs,
)

router = APIRouter(prefix="/api/import", tags=["import"])

BASE_STATUSES = ["установлен", "склад", "ремонт", "списан"]


def read_sheet(data):
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)

    ws = None

    for sheet in wb.worksheets:
        if sheet.title.lower() == "list":
            ws = sheet
            break

    if ws is None:
        for sheet in wb.worksheets:
            if sheet.title.lower() != "@":
                ws = sheet
                break

    if ws is None:
        wb.close()
        raise HTTPException(
            status_code=400,
            detail="В файле нет листа с данными.",
        )

    rows = list(ws.iter_rows(values_only=True))
    sheet_title = ws.title
    wb.close()

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="Лист с данными пуст.",
        )

    header_index = detect_header(rows)
    header_row = rows[header_index]

    col_map = {}
    unknown_columns = []

    for index, raw_header in enumerate(header_row):
        header = normalize_header(raw_header)

        if header is None:
            continue

        mapped = HEADER_MAPPING.get(header)

        if mapped is None:
            unknown_columns.append(cell_to_text(raw_header))
        else:
            col_map[index] = mapped

    return rows, header_index, col_map, unknown_columns, sheet_title


def get_or_create_location(session, parent_id, kind, name, code=None):
    if name is None and code is None:
        return None

    if name is None:
        name = code

    base_filters = [Location.kind == kind]

    if parent_id is None:
        base_filters.append(Location.parent_id.is_(None))
    else:
        base_filters.append(Location.parent_id == parent_id)

    find_filters = base_filters.copy()

    if code is not None:
        find_filters.append(Location.code == code)
    else:
        find_filters.append(Location.name == name)

    location = session.query(Location).filter(*find_filters).first()

    if location:
        return location

    max_sort = session.query(func.max(Location.sort)).filter(*base_filters).scalar() or 0

    location = Location(
        parent_id=parent_id,
        kind=kind,
        name=name,
        code=code,
        sort=max_sort + 1,
    )

    session.add(location)
    session.flush()

    return location


def get_or_create_person(session, full_name):
    person = (
        session.query(Person)
        .filter(func.lower(Person.full_name) == full_name.lower())
        .first()
    )

    if person:
        return person

    person = Person(full_name=full_name)

    session.add(person)
    session.flush()

    return person


def get_or_create_vacuum_account(session, login):
    account = (
        session.query(VacuumAccount)
        .filter(VacuumAccount.login == login)
        .first()
    )

    if account:
        return account

    account = VacuumAccount(login=login)

    session.add(account)
    session.flush()

    return account


def ensure_choice(session, field, value):
    value = clean_text(value)

    if not value:
        return

    exists = (
        session.query(Choice)
        .filter(Choice.field == field, Choice.value == value)
        .first()
    )

    if exists:
        return

    max_sort = session.query(func.max(Choice.sort)).filter(Choice.field == field).scalar() or 0

    session.add(
        Choice(
            field=field,
            value=value,
            sort=max_sort + 1,
        )
    )

    session.flush()


def seed_base_statuses(session):
    for index, value in enumerate(BASE_STATUSES):
        exists = (
            session.query(Choice)
            .filter(Choice.field == "status", Choice.value == value)
            .first()
        )

        if exists:
            continue

        session.add(
            Choice(
                field="status",
                value=value,
                sort=index,
            )
        )

    session.flush()


@router.post("/apply")
async def apply_import(
    file: UploadFile = File(...),
    confirm: bool = Form(False),
):
    name = (file.filename or "").lower()

    if not name.endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=400,
            detail="Нужен файл .xlsx или .xlsm. Если файл .xls — пересохрани его в .xlsx через Excel.",
        )

    data = await file.read()

    rows, header_index, col_map, unknown_columns, sheet_title = read_sheet(data)

    if unknown_columns:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "В файле есть неизвестные столбцы. Импорт без решения по ним не выполняется.",
                "unknown_columns": unknown_columns,
            },
        )

    if not confirm:
        return {
            "message": "Для записи в базу отправьте запрос ещё раз с confirm=true.",
            "sheet": sheet_title,
        }

    session = SessionLocal()

    try:
        existing_computers = session.query(func.count(Computer.id)).scalar()

        if existing_computers:
            raise HTTPException(
                status_code=409,
                detail="В базе уже есть компьютеры. Повторный импорт пока не поддерживается.",
            )

        seed_base_statuses(session)

        stats = {
            "rows": 0,
            "computers": 0,
            "anomalies": 0,
        }

        anomalies = []

        def add_anomaly(row_label, field, value, message):
            stats["anomalies"] += 1

            if len(anomalies) < 100:
                anomalies.append(
                    {
                        "row": row_label,
                        "field": field,
                        "value": value,
                        "message": message,
                    }
                )

        for data_index, row in enumerate(rows[header_index + 1 :]):
            if all(v is None for v in row):
                continue

            rec = {}

            for index, key in col_map.items():
                rec[key] = row[index] if index < len(row) else None

            if all(cell_to_text(rec.get(field)) is None for field in IMPORTANT_FIELDS):
                continue

            stats["rows"] += 1

            excel_row = header_index + data_index + 2
            row_num_text = cell_to_text(rec.get("row_num"))

            if row_num_text:
                row_label = f"строка {excel_row}, #{row_num_text}"
            else:
                row_label = f"строка {excel_row}"

            location_id = None

            address = clean_text(rec.get("address"))
            department = clean_text(rec.get("department"))

            floor_raw = clean_text(rec.get("floor"))
            floor = None

            if floor_raw:
                floor_int, floor_ok = parse_int(floor_raw)

                if floor_ok and floor_int is not None:
                    floor = str(floor_int)
                else:
                    floor = floor_raw

            room_code = clean_text(rec.get("room_code"))
            room_name = clean_text(rec.get("room_name"))

            if not address:
                add_anomaly(row_label, "Адрес", address, "Пустой адрес")

            if address:
                building = get_or_create_location(session, None, "building", address, None)
                location_id = building.id

                if department:
                    department_location = get_or_create_location(
                        session,
                        building.id,
                        "department",
                        department,
                        None,
                    )

                    location_id = department_location.id
                    parent_for_room = department_location.id

                    if floor:
                        floor_location = get_or_create_location(
                            session,
                            department_location.id,
                            "floor",
                            floor,
                            None,
                        )

                        location_id = floor_location.id
                        parent_for_room = floor_location.id

                    if room_code or room_name:
                        room_location = get_or_create_location(
                            session,
                            parent_for_room,
                            "room",
                            room_name,
                            room_code,
                        )

                        location_id = room_location.id
                else:
                    add_anomaly(row_label, "Отделение", department, "Пустое отделение")

            person = None
            person_name = clean_text(rec.get("person"))

            if person_name:
                person = get_or_create_person(session, person_name)

            seat_no, seat_ok = parse_int(rec.get("seat_no"))

            if not seat_ok:
                add_anomaly(
                    row_label,
                    "№",
                    cell_to_text(rec.get("seat_no")),
                    "Номер места не является целым числом",
                )
                seat_no = None

            row_num_int, row_num_ok = parse_int(rec.get("row_num"))

            if row_num_ok and row_num_int is not None:
                seat_sort = float(row_num_int)
            else:
                seat_sort = float(excel_row)

            ips, ip_issues = parse_ips(rec.get("ip"))

            for bad_ip, message in ip_issues:
                add_anomaly(row_label, "IP", bad_ip, message)

            macs, mac_issues = parse_macs(rec.get("mac"))

            for bad_mac, message in mac_issues:
                add_anomaly(row_label, "MAC", bad_mac, message)

            drive_values, drive_issues = parse_drive_text(rec.get("drive"))

            for message in drive_issues:
                add_anomaly(
                    row_label,
                    "DRIVE",
                    clean_text(rec.get("drive")),
                    message,
                )

            note = clean_text(rec.get("note"))

            printer = clean_text(rec.get("printer_skip"))

            if printer:
                add_anomaly(
                    row_label,
                    "Принтер",
                    printer,
                    "Значение из столбца Принтер пока не импортируется как принтер",
                )

                if note:
                    note = f"{note}\nПринтер: {printer}"
                else:
                    note = f"Принтер: {printer}"

            extra = {}

            for extra_key, source_key in (
                ("GSIT", "gsit"),
                ("Сост.", "state"),
                ("Метка", "label"),
            ):
                value = clean_text(rec.get(source_key))

                if value:
                    extra[extra_key] = value

            computer = Computer(
                status="установлен",
                location_id=location_id,
                seat_no=seat_no,
                seat_sort=seat_sort,
                temp_note=None,
                hostname=clean_text(rec.get("hostname")),
                ip="\n".join(ips) if ips else None,
                mac="\n".join(macs) if macs else None,
                inv_no=clean_text(rec.get("inv_no")),
                serial=None,
                type=clean_text(rec.get("type")),
                model=clean_text(rec.get("model")),
                os=clean_text(rec.get("os")),
                cpu=clean_text(rec.get("cpu")),
                ram=clean_text(rec.get("ram")),
                drive="\n".join(drive_values) if drive_values else None,
                gpu=clean_text(rec.get("gpu")),
                vnc=None,
                glpi_id=None,
                extra=extra,
                note=note,
                version=1,
            )

            session.add(computer)
            session.flush()

            stats["computers"] += 1

            for field, value in (
                ("os", computer.os),
                ("type", computer.type),
                ("model", computer.model),
                ("cpu", computer.cpu),
                ("gpu", computer.gpu),
            ):
                ensure_choice(session, field, value)

            for drive_value in drive_values:
                ensure_choice(session, "drive", drive_value)

            extra_field_map = {
                "GSIT": "gsit",
                "Сост.": "state",
                "Метка": "label",
            }
            
            for extra_key, extra_value in extra.items():
                field_name = extra_field_map.get(extra_key, extra_key)
                ensure_choice(session, field_name, extra_value)

            if person:
                link = (
                    session.query(ComputerPerson)
                    .filter(
                        ComputerPerson.computer_id == computer.id,
                        ComputerPerson.person_id == person.id,
                    )
                    .first()
                )

                if not link:
                    session.add(
                        ComputerPerson(
                            computer_id=computer.id,
                            person_id=person.id,
                            is_main=True,
                            sort=0,
                        )
                    )
                    session.flush()

            vacuum_raw = clean_text(rec.get("vacuum"))

            if vacuum_raw:
                for login in re.split(r"[\n\r,;]+", vacuum_raw):
                    login = login.strip().lower()

                    if not login:
                        continue

                    account = get_or_create_vacuum_account(session, login)

                    computer_link = (
                        session.query(VacuumAccountComputer)
                        .filter(
                            VacuumAccountComputer.account_id == account.id,
                            VacuumAccountComputer.computer_id == computer.id,
                        )
                        .first()
                    )

                    if not computer_link:
                        session.add(
                            VacuumAccountComputer(
                                account_id=account.id,
                                computer_id=computer.id,
                            )
                        )
                        session.flush()

                    if person:
                        person_link = (
                            session.query(VacuumAccountPerson)
                            .filter(
                                VacuumAccountPerson.account_id == account.id,
                                VacuumAccountPerson.person_id == person.id,
                            )
                            .first()
                        )

                        if not person_link:
                            session.add(
                                VacuumAccountPerson(
                                    account_id=account.id,
                                    person_id=person.id,
                                )
                            )
                            session.flush()

        session.commit()

        return {
            "ok": True,
            "sheet": sheet_title,
            "stats": stats,
            "anomalies": anomalies,
        }

    except HTTPException:
        session.rollback()
        raise

    except Exception as e:
        session.rollback()
        raise HTTPException(
            status_code=400,
            detail=f"Импорт не выполнен: {e}",
        )

    finally:
        session.close()