from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
import ipaddress
import re

from fastapi import APIRouter, File, HTTPException, UploadFile
from openpyxl import load_workbook

router = APIRouter(prefix="/api/import", tags=["import"])

HEADER_MAPPING = {
    "#": "row_num",
    "фио": "person",
    "адрес": "address",
    "отделение": "department",
    "эт.": "floor",
    "каб": "room_code",
    "кабинет": "room_name",
    "№": "seat_no",
    "ip": "ip",
    "hostname": "hostname",
    "vacuum": "vacuum",
    "os": "os",
    "тип": "type",
    "модель": "model",
    "cpu": "cpu",
    "ram": "ram",
    "drive": "drive",
    "gpu": "gpu",
    "mac": "mac",
    "инв": "inv_no",
    "gsit": "gsit",
    "сост.": "state",
    "метка": "label",
    "принтер": "printer_skip",
    "пр": "pr_skip",
    "#примечание": "note",
}

IMPORTANT_FIELDS = (
    "person",
    "address",
    "department",
    "room_code",
    "room_name",
    "hostname",
    "ip",
    "mac",
    "inv_no",
)

MAC_RE = re.compile(r"^([0-9A-F]{2}:){5}[0-9A-F]{2}$")

DRIVE_TYPED_RE = re.compile(
    r"(?P<type>SSD|HDD|NVME)\s*[-_]?\s*(?P<size>\d+)\s*(?P<unit>TB|GB|ТБ|ГБ|Т|Г)?",
    re.I,
)

DRIVE_REVERSE_RE = re.compile(
    r"(?P<size>\d+)\s*(?P<unit>TB|GB|ТБ|ГБ|Т|Г)?\s*(?P<type>SSD|HDD|NVME)",
    re.I,
)

BARE_DRIVE_RE = re.compile(
    r"\b(?P<size>\d+)\s*(?P<unit>TB|GB|ТБ|ГБ|Т|Г)?\b",
    re.I,
)


def cell_to_text(value):
    if value is None:
        return None

    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    text = str(value).strip()
    return text if text else None


def normalize_header(value):
    text = cell_to_text(value)
    if text is None:
        return None
    return text.lower()


def detect_header(rows):
    best_index = 0
    best_score = -1

    for index, row in enumerate(rows[:10]):
        headers = []

        for value in row:
            header = normalize_header(value)
            if header:
                headers.append(header)

        score = len(headers)

        if any(h in headers for h in ("ip", "hostname", "адрес", "фио")):
            score += 20

        if score > best_score:
            best_score = score
            best_index = index

    return best_index


def parse_int(value):
    text = cell_to_text(value)

    if text is None:
        return None, True

    try:
        number = float(text.replace(",", "."))
    except ValueError:
        return None, False

    if number.is_integer():
        return int(number), True

    return None, False


def clean_text(value):
    text = cell_to_text(value)

    if text is None:
        return None

    return " ".join(text.split())


def clean_multiline(value):
    """Как clean_text, но сохраняет переносы строк (для примечаний)."""
    text = cell_to_text(value)

    if text is None:
        return None

    lines = []

    for line in text.splitlines():
        line = " ".join(line.split())

        if line:
            lines.append(line)

    return "\n".join(lines) if lines else None


def split_vacuum_logins(value):
    """Логины VACUUM из ячейки: через перенос строки, пробел, запятую или ';'.

    Возвращает список в нижнем регистре без повторов (порядок сохраняется).
    """
    text = cell_to_text(value)

    if text is None:
        return []

    logins = []

    for login in re.split(r"[\s,;]+", text):
        login = login.strip().lower()

        if login and login not in logins:
            logins.append(login)

    return logins


def parse_ips(value):
    text = clean_text(value)

    if not text:
        return [], []

    ips = []
    issues = []

    for part in re.split(r"[\n\r,;\s]+", text):
        part = part.strip()

        if not part:
            continue

        try:
            ipaddress.ip_address(part)
            ips.append(part)
        except ValueError:
            issues.append((part, "IP не распознаётся"))

    return ips, issues


def parse_macs(value):
    text = cell_to_text(value)

    if not text:
        return [], []

    macs = []
    issues = []

    for part in re.split(r"[\n\r,;\s]+", text):
        part = part.strip().upper()

        if not part:
            continue

        if MAC_RE.match(part):
            macs.append(part)
        else:
            issues.append((part, "Неправильное значение MAC"))

    return macs, issues


def normalize_drive_size(size, unit):
    if unit:
        unit_upper = unit.upper()

        if unit_upper in ("TB", "ТБ", "Т"):
            return f"{size}TB"

    return str(size)


def parse_drive_text(value):
    text = clean_text(value)

    if text is None:
        return [], []

    upper = " ".join(text.upper().split())

    if not upper:
        return [], []

    values = []
    issues = []

    def remove_matches(source, regex, typed):
        parts = []
        last_end = 0

        for match in regex.finditer(source):
            size = match.group("size")
            unit = match.group("unit")

            if typed:
                drive_type = match.group("type").upper()
                values.append(f"{drive_type} {normalize_drive_size(size, unit)}")
            else:
                normalized = normalize_drive_size(size, unit)
                values.append(normalized)
                issues.append(f"Указан только объём накопителя: {normalized}")

            parts.append(source[last_end:match.start()])
            last_end = match.end()

        parts.append(source[last_end:])
        return " ".join(parts)

    remaining = remove_matches(upper, DRIVE_TYPED_RE, True)
    remaining = remove_matches(remaining, DRIVE_REVERSE_RE, True)
    remaining = remove_matches(remaining, BARE_DRIVE_RE, False)

    leftover = re.sub(r"[^0-9A-ZА-ЯЁ]+", " ", remaining, flags=re.I).strip()

    if not values and upper:
        values.append(upper)
        issues.append("Не удалось разобрать значение DRIVE")
    elif leftover:
        issues.append(f"Лишняя часть в DRIVE: {leftover}")

    return values, issues


def pack_duplicates(counter):
    result = {}

    for value, rows in counter.items():
        unique_rows = []

        for row in rows:
            if row not in unique_rows:
                unique_rows.append(row)

        if len(unique_rows) > 1:
            result[value] = unique_rows

    return result


# Поля строки файла, которые импорт кладёт в польз. поля ПК: ключ в extra → поле строки
EXTRA_FIELDS = (("GSIT", "gsit"), ("Сост.", "state"), ("Метка", "label"))
# Справочники из файла (поле строки = поле ПК); диски – отдельно (их в ячейке несколько)
CHOICE_FIELDS = ("os", "type", "model", "cpu", "gpu")


def check_file_name(file):
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=400,
            detail="Нужен файл .xlsx или .xlsm. Если файл .xls – пересохрани его в .xlsx через Excel.",
        )


def open_workbook(data):
    try:
        return load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Не удалось открыть файл: {e}",
        )


def read_sheet(data):
    """Лист с данными (list, иначе первый не «@») → (строки, номер строки заголовков,
    {номер столбца: поле}, неизвестные столбцы, название листа)."""
    wb = open_workbook(data)

    ws = None

    for sheet in wb.worksheets:
        if sheet.title.lower() == "list":
            ws = sheet
            break

    if ws is None:
        for sheet in wb.worksheets:
            if sheet.title.lower() != "@":
                ws = sheet
                break

    if ws is None:
        wb.close()
        raise HTTPException(
            status_code=400,
            detail="В файле нет листа с данными.",
        )

    rows = list(ws.iter_rows(values_only=True))
    sheet_title = ws.title
    wb.close()

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="Лист с данными пуст.",
        )

    header_index = detect_header(rows)
    header_row = rows[header_index]

    col_map = {}
    unknown_columns = []

    for index, raw_header in enumerate(header_row):
        header = normalize_header(raw_header)

        if header is None:
            continue

        mapped = HEADER_MAPPING.get(header)

        if mapped is None:
            unknown_columns.append(cell_to_text(raw_header))
        else:
            col_map[index] = mapped

    return rows, header_index, col_map, unknown_columns, sheet_title


def sheet_records(rows, header_index, col_map):
    """Строки данных: (номер строки в Excel, {поле: значение}); у строки без единого
    важного поля вместо словаря – None. Совсем пустые строки пропускаются."""
    for data_index, row in enumerate(rows[header_index + 1 :]):
        if all(v is None for v in row):
            continue

        rec = {key: row[index] if index < len(row) else None for index, key in col_map.items()}
        excel_row = header_index + data_index + 2

        if all(cell_to_text(rec.get(field)) is None for field in IMPORTANT_FIELDS):
            yield excel_row, None
        else:
            yield excel_row, rec


def row_label_of(rec, excel_row):
    """«строка 12, #10» – как строку называть в замечаниях."""
    row_num_text = cell_to_text(rec.get("row_num"))

    if row_num_text:
        return f"строка {excel_row}, #{row_num_text}"

    return f"строка {excel_row}"


def place_of(rec):
    """Расположение из строки: (адрес, отделение, этаж, код кабинета, название кабинета);
    этаж-число – без дробной части."""
    floor_raw = clean_text(rec.get("floor"))
    floor = None

    if floor_raw:
        floor_int, floor_ok = parse_int(floor_raw)

        if floor_ok and floor_int is not None:
            floor = str(floor_int)
        else:
            floor = floor_raw

    return (
        clean_text(rec.get("address")),
        clean_text(rec.get("department")),
        floor,
        clean_text(rec.get("room_code")),
        clean_text(rec.get("room_name")),
    )


class Anomalies:
    """Замечания к строкам файла: считаются все, хранятся первые limit (None – все)."""

    def __init__(self, limit=None):
        self.items = []
        self.count = 0
        self.limit = limit

    def add(self, row_label, field, value, message):
        self.count += 1

        if self.limit is None or len(self.items) < self.limit:
            self.items.append(
                {
                    "row": row_label,
                    "field": field,
                    "value": value,
                    "message": message,
                }
            )


@router.post("/preview")
async def preview(file: UploadFile = File(...)):
    check_file_name(file)
    wb = open_workbook(await file.read())

    sheets = []

    for ws in wb.worksheets:
        rows = []

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 20:
                break

            rows.append([cell_to_text(v) for v in row])

        sheets.append(
            {
                "name": ws.title,
                "rows": rows,
            }
        )

    wb.close()

    return {
        "filename": file.filename,
        "sheets": sheets,
    }


class Analysis:
    """Что импорт сделает с файлом (analyze): копится по строкам (add), итог – report."""

    def __init__(self):
        self.anomalies = Anomalies()
        self.total_rows = 0
        self.skipped_empty_rows = 0
        self.duplicates = {key: defaultdict(list) for key in ("ip", "mac", "hostname", "inv_no", "vacuum")}
        self.locations = {}
        self.users = set()
        self.vacuum_logins = set()
        self.choices = {field: set() for field in CHOICE_FIELDS + ("drive",)}
        self.extra_fields = {key: set() for key, _ in EXTRA_FIELDS}
        self.printer_values = set()
        self.pr_values = set()

    def add(self, rec, row_label):
        add_anomaly = self.anomalies.add
        address, department, floor, room_code, room_name = place_of(rec)

        if not address:
            add_anomaly(row_label, "Адрес", address, "Пустой адрес")

        if not department:
            add_anomaly(row_label, "Отделение", department, "Пустое отделение")

        if room_code and room_name:
            room_display = f"{room_code} {room_name}"
        else:
            room_display = room_name or room_code

        path_parts = [p for p in [address, department, floor, room_display] if p]

        if path_parts:
            path = " / ".join(path_parts)
            self.locations[path] = self.locations.get(path, 0) + 1

        person = clean_text(rec.get("person"))
        if person:
            self.users.add(person)

        _, seat_ok = parse_int(rec.get("seat_no"))
        if not seat_ok:
            add_anomaly(
                row_label,
                "№",
                cell_to_text(rec.get("seat_no")),
                "Номер места не является целым числом",
            )

        ips, ip_issues = parse_ips(rec.get("ip"))

        for ip in ips:
            self.duplicates["ip"][ip].append(row_label)

        for bad_ip, message in ip_issues:
            add_anomaly(row_label, "IP", bad_ip, message)

        hostname = clean_text(rec.get("hostname"))
        if hostname:
            self.duplicates["hostname"][hostname.lower()].append(row_label)

        for login in split_vacuum_logins(rec.get("vacuum")):
            self.vacuum_logins.add(login)
            self.duplicates["vacuum"][login].append(row_label)

        for field in CHOICE_FIELDS:
            value = clean_text(rec.get(field))
            if value:
                self.choices[field].add(value)

        drive_values, drive_issues = parse_drive_text(rec.get("drive"))
        self.choices["drive"].update(drive_values)

        for message in drive_issues:
            add_anomaly(
                row_label,
                "DRIVE",
                clean_text(rec.get("drive")),
                message,
            )

        macs, mac_issues = parse_macs(rec.get("mac"))

        for mac in macs:
            self.duplicates["mac"][mac].append(row_label)

        for bad_mac, message in mac_issues:
            add_anomaly(row_label, "MAC", bad_mac, message)

        inv_no = clean_text(rec.get("inv_no"))
        if inv_no:
            self.duplicates["inv_no"][inv_no].append(row_label)

        for key, field in EXTRA_FIELDS:
            value = clean_text(rec.get(field))
            if value:
                self.extra_fields[key].add(value)

        printer = clean_text(rec.get("printer_skip"))
        if printer:
            self.printer_values.add(printer)
            add_anomaly(
                row_label,
                "Принтер",
                printer,
                "Значение из столбца Принтер пока не импортируется как принтер",
            )

        pr = clean_text(rec.get("pr_skip"))
        if pr:
            self.pr_values.add(pr)

    def report(self):
        locations_preview = [
            {"path": path, "count": count}
            for path, count in sorted(self.locations.items(), key=lambda item: item[0].lower())[:200]
        ]

        return {
            "summary": {
                "users": len(self.users),
                "vacuum_logins": len(self.vacuum_logins),
                "locations": len(self.locations),
            },
            "locations_preview": locations_preview,
            "users_preview": sorted(self.users, key=str.lower)[:200],
            "vacuum_preview": sorted(self.vacuum_logins, key=str.lower)[:200],
            "choices": {
                key: sorted(values, key=str.lower)
                for key, values in self.choices.items()
            },
            "extra_fields": {
                key: sorted(values, key=str.lower)
                for key, values in self.extra_fields.items()
            },
            "skipped_values": {
                "пр": sorted(self.pr_values, key=str.lower),
                "принтер": sorted(self.printer_values, key=str.lower),
            },
            "duplicates": {key: pack_duplicates(rows) for key, rows in self.duplicates.items()},
            "anomalies_count": self.anomalies.count,
            "anomalies": self.anomalies.items[:200],
        }


@router.post("/analyze")
async def analyze(file: UploadFile = File(...)):
    check_file_name(file)
    rows, header_index, col_map, unknown_columns, sheet_title = read_sheet(await file.read())
    analysis = Analysis()

    for excel_row, rec in sheet_records(rows, header_index, col_map):
        if rec is None:
            analysis.skipped_empty_rows += 1
            continue

        analysis.total_rows += 1
        analysis.add(rec, row_label_of(rec, excel_row))

    return {
        "filename": file.filename,
        "sheet": sheet_title,
        "header_row": header_index + 1,
        "headers": [cell_to_text(v) for v in rows[header_index]],
        "total_rows": analysis.total_rows,
        "skipped_empty_rows": analysis.skipped_empty_rows,
        "unknown_columns": unknown_columns,
        **analysis.report(),
    }
