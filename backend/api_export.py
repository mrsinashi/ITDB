from datetime import date
from io import BytesIO

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from api_computers import is_reserved_field_key, list_computers
from db import SessionLocal
from models import FieldDef

router = APIRouter(prefix="/api/export", tags=["export"])

COLUMNS = [
    {"field": "user", "header": "ФИО", "width": 25},
    {"field": "building", "header": "Адрес", "width": 16},
    {"field": "department", "header": "Отделение", "width": 14},
    {"field": "floor", "header": "Эт.", "width": 6},
    {"field": "room_code", "header": "Каб", "width": 8},
    {"field": "room_name", "header": "Кабинет", "width": 18},
    {"field": "seat_no", "header": "№", "width": 6},
    {"field": "ip", "header": "IP", "width": 16, "wrap": True},
    {"field": "hostname", "header": "HOSTNAME", "width": 18},
    {"field": "vacuum", "header": "VACUUM", "width": 16, "wrap": True},
    {"field": "os", "header": "OS", "width": 14},
    {"field": "type", "header": "ТИП", "width": 10},
    {"field": "model", "header": "Модель", "width": 16},
    {"field": "cpu", "header": "CPU", "width": 20},
    {"field": "ram", "header": "RAM", "width": 7},
    {"field": "drive", "header": "DRIVE", "width": 16, "wrap": True},
    {"field": "gpu", "header": "GPU", "width": 16},
    {"field": "mac", "header": "MAC", "width": 20, "wrap": True},
    {"field": "inv_no", "header": "ИНВ", "width": 10},
    {"field": "gsit", "header": "GSIT", "width": 8},
    {"field": "state", "header": "Сост.", "width": 8},
    {"field": "label", "header": "Метка", "width": 10},
    {"field": "status", "header": "Статус", "width": 10},
    {"field": "note", "header": "Примечание", "width": 30, "wrap": True},
]


def get_export_columns():
    """Встроенные столбцы + пользовательские поля (перед «Статус», как в таблице)."""
    session = SessionLocal()

    try:
        field_defs = (
            session.query(FieldDef)
            .filter(FieldDef.archived == False)
            .order_by(FieldDef.sort, FieldDef.id)
            .all()
        )

        extra_columns = [
            {"field": fd.key, "header": fd.label, "width": 14, "wrap": True}
            for fd in field_defs
            if not is_reserved_field_key(fd.key)
        ]
    finally:
        session.close()

    status_index = next(
        (i for i, column in enumerate(COLUMNS) if column["field"] == "status"),
        len(COLUMNS),
    )

    return COLUMNS[:status_index] + extra_columns + COLUMNS[status_index:]


@router.get("/computers.xlsx")
def export_computers():
    data = list_computers()
    rows = data.get("rows", [])
    columns = get_export_columns()

    wb = Workbook()
    ws = wb.active
    ws.title = "Компьютеры"

    header_font = Font(bold=True)
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    for col_index, column in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_index, value=column["header"])
        cell.font = header_font
        ws.column_dimensions[cell.column_letter].width = column["width"]

    for row_index, row in enumerate(rows, start=2):
        for col_index, column in enumerate(columns, start=1):
            value = row.get(column["field"])
            cell = ws.cell(row=row_index, column=col_index, value=value)

            if column.get("wrap"):
                cell.alignment = wrap_alignment

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"itdb_computers_{date.today().isoformat()}.xlsx"

    return StreamingResponse(
        buffer,
        headers={
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "Content-Disposition": f"attachment; filename={filename}",
        },
    )