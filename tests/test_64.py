"""Этап 64: одни названия везде – столбцы (VACUUM, «№ места», «Тип», «Временно до») и
«Создан» в Истории у всех объектов."""
from io import BytesIO

from openpyxl import load_workbook

from conftest import ok


def by_key(columns):
    return {column["key"]: column for column in columns}


def test_column_names(reader):
    data = ok(reader.get("/api/columns"))
    pcs = by_key(data["columns"])
    printers = by_key(data["printer_columns"])

    # Вне таблицы – так же, как в шапке; подпись карточки – та же
    assert (pcs["vacuum"]["short"], pcs["vacuum"]["title"], pcs["vacuum"]["card"]) == ("VACUUM", "VACUUM", "VACUUM")
    assert (pcs["type"]["short"], pcs["type"]["title"]) == ("Тип", "Тип компьютера")
    assert (pcs["temp_until"]["title"], pcs["temp_until"]["card"]) == ("Временно до", "Временно до")
    assert pcs["seat_no"]["title"] == "№ места"
    assert pcs["room_code"]["title"] == printers["room_code"]["title"] == "№ кабинета"
    assert printers["number"]["title"] == "№ в кабинете"
    assert data["history_labels"]["temp_note"] == "Временно до"


def test_export_header_type(admin):
    ws = load_workbook(BytesIO(admin.get("/api/export/computers.xlsx").content))["Компьютеры"]
    head = [cell.value for cell in ws[1]]
    assert "Тип" in head and "ТИП" not in head and "VACUUM" in head


def test_created_label_everywhere(admin, editor):
    ok(admin.post("/api/users", json={"login": "u64", "role": "reader", "password": "secret1"}))
    ok(editor.post("/api/choices", json={"field": "os", "value": "Win 64"}))
    data = ok(admin.get("/api/history", params={"q": "создан", "tz": "Asia/Irkutsk"}))
    entities = {item["entity"] for item in data["items"]}
    assert {"users", "choices"} <= entities
    # «Создано» больше нигде не пишется
    assert not ok(admin.get("/api/history", params={"q": "создано", "tz": "Asia/Irkutsk"}))["items"]
