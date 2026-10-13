"""Этап 46: MAC принтера – столбец после «Доступа», прописными строками, в карточке после IP,
в выгрузке, с отменой; «Компьютеры» у принтеров – после IP."""
from io import BytesIO

from openpyxl import load_workbook

from conftest import history, ok
from test_printers import add_printer, patch, printer


def test_printer_columns_order(reader):
    info = ok(reader.get("/api/columns"))
    keys = [c["key"] for c in info["printer_columns"]]
    assert keys[keys.index("ip") + 1] == "computers"
    assert keys[keys.index("web_auth") + 1] == "mac"

    mac = next(c for c in info["printer_columns"] if c["key"] == "mac")
    assert (mac["short"], mac["kind"], mac["dup"], mac["bulk"], mac["card_copy"]) == ("MAC", "mac", True, False, True)

    network = next(g for g in info["printer_card_groups"] if g["title"] == "Сеть")
    assert network["fields"][:2] == ["ip", "mac"]


def test_printer_mac(editor, reader, room):
    pid = add_printer(editor, room["room"], name="prn-mac")
    patch(editor, pid, mac=" d8:bb:cc:dd:ee:01,  d8:bb:cc:dd:ee:02 ")
    assert printer(reader, pid)["mac"] == "D8:BB:CC:DD:EE:01\nD8:BB:CC:DD:EE:02"
    assert history(reader, "printers", pid)[0]["changes"]["mac"] == {"old": None, "new": "D8:BB:CC:DD:EE:01\nD8:BB:CC:DD:EE:02"}

    # У каждого свой – массово не меняется; читатель не правит
    assert editor.post("/api/printers/bulk-update", json={"ids": [pid], "field": "mac", "value": "x"}).status_code == 400
    assert reader.patch(f"/api/printers/{pid}", json={"mac": "x", "_version": printer(reader, pid)["version"]}).status_code == 403

    # Выгрузка: после «Web» («Доступ» не выгружается)
    wb = load_workbook(BytesIO(reader.get("/api/export/printers.xlsx").content))
    header = [cell.value for cell in wb["Принтеры"][1]]
    assert header[header.index("Web") + 1] == "MAC"
    assert wb["Принтеры"].cell(row=2, column=header.index("MAC") + 1).value == "D8:BB:CC:DD:EE:01\nD8:BB:CC:DD:EE:02"

    # Ctrl+Z
    ok(editor.post("/api/history/undo-last"))
    assert printer(reader, pid)["mac"] is None
