"""Этап 26е: объём дисков и MAC из GLPI / GSIT, имя ПК – только сообщить,
VACUUM из Jabber, свои названия антивирусов."""
from io import BytesIO

from openpyxl import load_workbook

import scan_normalize as sn
from conftest import get_row, ok
from scan_values import Names
from test_scan_collect import ETH, Glpi, add_pc, glpi_url, pc, reset_glpi, secret_key, setup_source  # noqa: F401
from test_scan_collect import collect as glpi_collect
from test_scan_jabber import DOMAIN, Ejabberd, jabber_url, reset_jabber, setup_jabber  # noqa: F401
from test_scan_jabber import collect as jabber_collect

GB = 1e9 / (1024 * 1024)   # 1 ГБ «с наклейки» в двоичных МБ, как пишет агент


def mib(gb):
    return int(gb * GB)


def diffs_of(client):
    return {(d["computer_id"], d["field"]): d for d in ok(client.get("/api/scan/diffs"))["items"]}


# ---------- Диски ----------


def test_disk_sizes_binary_mb():
    # Агент пишет объём в двоичных МБ: 128 ГБ – это 122 070, а не 128 000
    assert sn.disk_size_text(mib(128)) == "128"
    assert sn.disk_size_text(mib(256)) == "256"
    assert sn.disk_size_text(mib(240)) == "240"
    assert sn.disk_size_text(mib(480)) == "480"
    assert sn.disk_size_text(mib(512)) == "512"
    assert sn.disk_size_text(mib(500)) == "500"
    assert sn.disk_size_text(mib(1000)) == "1TB"
    # Объём из названия модели, если сходится
    assert sn.disk_size_text(mib(240), "KINGSTON SA400S37240G") == "240"
    assert sn.disk_size_text(mib(480), "CT480BX500SSD1") == "480"
    assert sn.disk_size_text(mib(256), "ADATA SU650 256GB") == "256"
    # Название говорит другое – не верим ему
    assert sn.disk_size_text(mib(500), "SSD 128GB") == "500"
    # Источник с обычными МБ – тоже понимается
    assert sn.disk_size_text(250059, "Samsung SSD 870 EVO 250GB") == "250"
    assert sn.disk_size_text(1000204) == "1TB"
    assert sn.drives_short([
        {"name": "KINGSTON SA400S37120G", "mb": mib(120)},
        {"name": "Netac SSD 128GB", "mb": mib(128)},
    ]) == "SSD 120\nSSD 128"


# ---------- MAC ----------


def test_mac_values_and_keys():
    values, keys, _ = sn.build({"ports": [
        {"name": "Realtek PCIe GbE Family Controller", "mac": "d8:bb:c1:00:00:01", "ips": ["10.0.2.11"]},
        {"name": "Intel(R) Wi-Fi 6 AX201 160MHz", "mac": "da:bb:c1:00:00:02", "ips": ["10.0.5.20"]},
        {"name": "Bluetooth Device (Personal Area Network)", "mac": "d8:bb:c1:00:00:03", "ips": []},
        {"name": "Microsoft Wi-Fi Direct Virtual Adapter", "mac": "da:bb:c1:00:00:04", "ips": []},
        {"name": "TAP-Windows Adapter V9", "mac": "00:ff:12:34:56:78", "ips": ["10.8.0.6"]},
        {"name": "VirtualBox Host-Only Ethernet Adapter", "mac": "0a:00:27:00:00:05", "ips": ["192.168.56.1"]},
    ]})
    # В столбец – все адаптеры железа (Wi-Fi со случайным MAC, Bluetooth), без программных
    assert values["mac"] == "D8:BB:C1:00:00:01\nDA:BB:C1:00:00:02\nD8:BB:C1:00:00:03"
    assert values["ip"] == "10.0.2.11\n10.0.5.20"
    # Признаки – только заводские MAC физических адаптеров (как раньше)
    assert keys["macs"] == ["D8:BB:C1:00:00:01"]

    # В таблице больше MAC, чем видит источник, – не расхождение; меньше – «≈»
    names = Names()
    three = "D8:BB:C1:00:00:01\nD8:BB:C1:00:00:02\nD8:BB:C1:00:00:03"
    assert names.compare("mac", "D8:BB:C1:00:00:01\nD8:BB:C1:00:00:03", three)["mark"] == "="
    assert names.compare("mac", three, "D8:BB:C1:00:00:01")["mark"] == "≈"
    assert names.compare("mac", "D8:BB:C1:00:00:09", three)["mark"] == "≠"
    # Общие MAC есть, но источник видит и новый – дописать, MAC таблицы не убирать
    both = names.compare("mac", "D8:BB:C1:00:00:01\nD8:BB:C1:00:00:09", three)
    assert both["mark"] == "≠" and both["source"] == three + "\nD8:BB:C1:00:00:09"
    assert both["raw"] == "D8:BB:C1:00:00:01\nD8:BB:C1:00:00:09"
    # Общих нет – предлагается, как видит источник
    assert names.compare("mac", "D8:BB:C1:00:00:09", three)["source"] == "D8:BB:C1:00:00:09"


def test_network_cards_without_ports():
    """MAC сетевой карты из устройств, у которой нет порта, – тоже в столбец."""
    from scan_glpi import computer_details

    class Fake:
        def call(self, path, **params):
            import json

            class R:
                status = 200
                body = json.dumps({
                    "name": "pc", "_devices": {"Item_DeviceNetworkCard": {
                        "1": {"devicenetworkcards_id": "Realtek PCIe GbE", "mac": "d8:bb:c1:00:00:01"},
                        "2": {"devicenetworkcards_id": "Intel Wireless-AC 9560", "mac": "d8:bb:c1:00:00:02"},
                    }},
                    "_networkports": {"NetworkPortEthernet": [{"name": "Ethernet", "mac": "d8:bb:c1:00:00:01"}]},
                }).encode()
            return R()

    raw = computer_details(Fake(), 1)
    assert [p["mac"] for p in raw["ports"]] == ["d8:bb:c1:00:00:01", "d8:bb:c1:00:00:02"]


# ---------- HOSTNAME – только сообщить ----------


def test_hostname_only_reported(admin, editor, room, glpi_url):
    loc = room["room"]
    named = add_pc(editor, loc, "ter-201-1", mac="D8:BB:C1:00:00:01")
    empty = add_pc(editor, loc, "", mac="D8:BB:C1:00:00:02")
    Glpi.computers = {
        1: pc("DESKTOP-AB12CD", ports=[(ETH, "d8:bb:c1:00:00:01", [])]),
        2: pc("reg-101-1", ports=[(ETH, "d8:bb:c1:00:00:02", [])]),
    }
    setup_source(admin, glpi_url)
    assert glpi_collect(admin)["status"] == "ok"

    diffs = diffs_of(editor)
    d = diffs[(named, "hostname")]
    assert (d["kind"], d["proposed"], d["can_take"]) == ("diff", "DESKTOP-AB12CD", False)
    assert diffs[(empty, "hostname")]["can_take"] is True
    assert diffs[(named, "model")]["can_take"] is True

    # «Взять» имя не записывает
    result = ok(editor.post("/api/scan/diffs/accept", json={"items": [
        {"computer_id": named, "field": "hostname", "value": "DESKTOP-AB12CD", "table": "ter-201-1"},
        {"computer_id": empty, "field": "hostname", "value": "reg-101-1", "table": ""},
    ]}))
    assert result["accepted"] == 1 and len(result["skipped"]) == 1
    assert get_row(editor, named)["hostname"] == "ter-201-1"
    assert get_row(editor, empty)["hostname"] == "reg-101-1"

    # «Не показывать» – как отклонение
    ok(editor.post("/api/scan/diffs/reject", json={"items": [{"computer_id": named, "field": "hostname", "raw": "DESKTOP-AB12CD"}]}))
    assert (named, "hostname") not in diffs_of(editor)


# ---------- VACUUM и IP из Jabber ----------


def set_vacuum(client, computer_id, logins):
    ok(client.patch(f"/api/computers/{computer_id}", json={"vacuum": logins, "_version": get_row(client, computer_id)["version"]}))


def test_jabber_vacuum(admin, editor, room, glpi_url, jabber_url):
    """VACUUM из Jabber – только по IP ПК (этап 26ж: IP по логину больше не предлагается)."""
    loc = room["room"]
    empty_vac = add_pc(editor, loc, "pc-empty", ip="10.0.2.11")
    has_other = add_pc(editor, loc, "pc-other", ip="10.0.2.12")
    set_vacuum(editor, has_other, "sidorov")
    has_all = add_pc(editor, loc, "pc-all", ip="10.0.2.13")
    set_vacuum(editor, has_all, "kuznetsova\nsmirnov")
    visitor_pc = add_pc(editor, loc, "pc-visit", ip="10.0.2.14")
    no_ip = add_pc(editor, loc, "pc-noip")
    set_vacuum(editor, no_ip, "novikova")
    glpi_ip = add_pc(editor, loc, "pc-glpi", mac="D8:BB:C1:00:00:07")
    dup_a = add_pc(editor, loc, "pc-dup-a", ip="10.0.2.20")
    add_pc(editor, loc, "pc-dup-b", ip="10.0.2.20")

    Glpi.computers = {7: pc("pc-glpi", ports=[(ETH, "d8:bb:c1:00:00:07", ["10.0.2.77"])])}
    setup_source(admin, glpi_url)
    assert glpi_collect(admin)["status"] == "ok"

    Ejabberd.online = {
        "ivanov": [("Vacuum-IM", "10.0.2.11")],          # ПК без VACUUM → заполнить
        "petrova": [("Vacuum-IM", "10.0.2.12")],         # у ПК sidorov → добавить
        "kuznetsova": [("Vacuum-IM", "10.0.2.13")],      # уже записана → ничего
        "sidorov": [("Vacuum-IM", "10.0.2.14")],         # записан у другого ПК → неточно
        "novikova": [("Vacuum-IM", "10.0.4.40")],        # по логину ПК не ищется: IP не предлагается
        "vasiliev": [("Vacuum-IM", "10.0.2.77")],        # адрес ПК знает только GLPI → VACUUM этому ПК
        "orlov": [("Vacuum-IM", "10.0.2.20")],           # адрес у двух ПК – не понять
    }
    setup_jabber(admin, jabber_url)
    assert jabber_collect(admin)["status"] == "ok"

    data = ok(editor.get("/api/scan/diffs"))
    diffs = {(d["computer_id"], d["field"]): d for d in data["items"]}
    d = diffs[(empty_vac, "vacuum")]
    assert (d["kind"], d["table"], d["proposed"], d["raw"]) == ("fill", "", "ivanov", "ivanov")
    assert d["sources"][0]["source"] == "jabber" and d["sources"][0]["source_id"] is None and d["sources"][0]["by"] == ["ip"]
    assert "ivanov" in d["note"]
    d = diffs[(has_other, "vacuum")]
    assert (d["kind"], d["table"], d["proposed"]) == ("diff", "sidorov", "petrova\nsidorov")
    assert (has_all, "vacuum") not in diffs
    d = diffs[(visitor_pc, "vacuum")]
    assert d["kind"] == "unsure" and "pc-other" in d["unsure"]
    assert (no_ip, "ip") not in diffs and (no_ip, "vacuum") not in diffs
    assert diffs[(glpi_ip, "vacuum")]["proposed"] == "vasiliev"
    assert diffs[(glpi_ip, "ip")]["sources"][0]["source"] == "glpi"
    assert (dup_a, "vacuum") not in diffs
    # Кого Jabber видит с адреса ПК – для столбца «Vacuum» в подробностях «Проверки»
    assert data["jabber"][str(has_all)] == ["kuznetsova"] and data["jabber"][str(has_other)] == ["petrova"]
    # Списка пользователей нет – «нет в Jabber» не помечается
    assert data["vacuum_missing"] == []

    # Взять VACUUM – дописывается к тому, что было
    result = ok(editor.post("/api/scan/diffs/accept", json={"items": [
        {"computer_id": has_other, "field": "vacuum", "value": "petrova\nsidorov", "table": "sidorov", "source": "Jabber"},
    ]}))
    assert result == {"accepted": 1, "skipped": []}
    assert get_row(editor, has_other)["vacuum"] == "petrova\nsidorov"
    history = ok(editor.get("/api/history", params={"entity": "computers", "entity_id": has_other}))["items"]
    assert history[0]["changes"]["vacuum"]["scan"] == "Jabber"

    # Таблица изменилась, пока смотрели, – не пишется
    result = ok(editor.post("/api/scan/diffs/accept", json={"items": [
        {"computer_id": empty_vac, "field": "vacuum", "value": "ivanov", "table": "someone"},
    ]}))
    assert result["accepted"] == 0

    # Отклонить VACUUM – больше не предлагается
    ok(editor.post("/api/scan/diffs/reject", json={"items": [{"computer_id": empty_vac, "field": "vacuum", "raw": "ivanov"}]}))
    assert (empty_vac, "vacuum") not in diffs_of(editor)

    # Jabber выключен – не предлагает
    ok(admin.patch("/api/scan/sources/jabber", json={"enabled": False}))
    diffs = diffs_of(editor)
    assert not any(d["sources"][0]["source"] == "jabber" for d in diffs.values())


# ---------- Свои названия антивирусов ----------


def test_antivirus_rename(admin, editor, reader, room, glpi_url):
    loc = room["room"]
    first = add_pc(editor, loc, "pc-1", mac="D8:BB:C1:00:00:01")
    Glpi.computers = {1: pc("pc-1", ports=[(ETH, "d8:bb:c1:00:00:01", [])])}
    long_name = "Kaspersky Endpoint Security для Windows"
    Glpi.antivirus = [{"id": 1, "computers_id": 1, "name": long_name, "is_active": 1, "is_uptodate": 1}]
    setup_source(admin, glpi_url)
    assert glpi_collect(admin)["status"] == "ok"

    # Начальный вид – блочок с фоном по состоянию
    settings = ok(reader.get("/api/scan/antivirus"))
    assert all(s["bg_color"] and not s["color"] for s in settings["statuses"]) and settings["names"] == {}

    assert reader.patch("/api/scan/antivirus", json={"rename": {"source": long_name, "name": "KES"}}).status_code == 403
    assert editor.patch("/api/scan/antivirus", json={"rename": {"source": " ", "name": "KES"}}).status_code == 400
    assert editor.patch("/api/scan/antivirus", json={"rename": {"source": long_name, "name": "K" * 61}}).status_code == 400
    settings = ok(editor.patch("/api/scan/antivirus", json={"rename": {"source": long_name, "name": " KES "}}))
    assert settings["names"] == {long_name.lower(): "KES"}
    assert ok(reader.get("/api/scan/antivirus"))["names"] == {long_name.lower(): "KES"}

    # Скрытые названия и названия по-прежнему – как в источнике
    av = ok(reader.get("/api/scan/diffs"))["antivirus"]
    assert av[str(first)][0]["name"] == long_name

    # Выгрузка – своим названием
    response = reader.get("/api/export/computers.xlsx")
    ws = load_workbook(BytesIO(response.content))["Компьютеры"]
    headers = [cell.value for cell in ws[1]]
    assert ws.cell(2, headers.index("Антивирус") + 1).value == "KES"

    # История: было – стало
    items = ok(reader.get("/api/history", params={"entity": "scan_antivirus"}))["items"]
    assert items[0]["title"] == long_name and items[0]["changes"]["name"] == {"old": long_name, "new": "KES"}

    # Пустое название – снова как в источнике
    settings = ok(editor.patch("/api/scan/antivirus", json={"rename": {"source": long_name, "name": ""}}))
    assert settings["names"] == {}
