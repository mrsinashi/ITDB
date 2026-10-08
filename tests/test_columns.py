"""Единое описание столбцов (GET /api/columns) и поля этапа 16:
серийный номер, VNC, GLPI ID, «Временно, до…» (дата)."""
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from api_columns import COLUMNS, CARD_GROUPS
from api_computers import RESERVED_FIELD_KEYS
from conftest import add_computer, get_row, ok

NEW_FIELDS = {
    "serial": "SN-001",
    "vnc": "tight",
    "glpi_id": "1234",
    "temp_until": "15.10.2026",
}


def test_columns_for_any_role(reader, anon):
    data = ok(reader.get("/api/columns"))
    keys = [column["key"] for column in data["columns"]]

    assert keys == [column.key for column in COLUMNS]
    assert {"serial", "vnc", "glpi_id", "temp_until"} <= set(keys)
    assert data["history_labels"]["location_id"] == "Расположение"

    serial = next(column for column in data["columns"] if column["key"] == "serial")
    assert serial["hidden"] and serial["dup"] and not serial["bulk"]

    assert anon.get("/api/columns").status_code == 401


def test_description_is_consistent():
    keys = {column.key for column in COLUMNS}

    assert len(keys) == len(COLUMNS), "ключ столбца повторяется"
    assert keys <= RESERVED_FIELD_KEYS

    for _, fields in CARD_GROUPS:
        for field in fields:
            assert field == "user_fields" or field in keys, field


def test_row_has_every_column(editor, room):
    computer_id = add_computer(editor, room["room"], hostname="pc-1")
    row = get_row(editor, computer_id)

    for column in COLUMNS:
        # Столбцы из сканера (антивирусы) в строке ПК не хранятся – их дописывает фронт
        assert column.key in row or column.kind == "scan", column.key


def test_new_fields_edit_and_history(editor, room):
    computer_id = add_computer(editor, room["room"], hostname="pc-1")
    row = get_row(editor, computer_id)

    payload = {key: f"  {value}  " for key, value in NEW_FIELDS.items()}
    payload["_version"] = row["version"]
    data = ok(editor.patch(f"/api/computers/{computer_id}", json=payload))

    for key, value in NEW_FIELDS.items():
        assert data["updated"][key] == value
        assert data["changes"][key] == {"old": None, "new": value}

    row = get_row(editor, computer_id)
    for key, value in NEW_FIELDS.items():
        assert row[key] == value

    # Очистить
    data = ok(editor.patch(f"/api/computers/{computer_id}", json={"temp_until": ""}))
    assert data["updated"]["temp_until"] is None
    assert get_row(editor, computer_id)["temp_until"] is None


def test_duplicate_serial_is_saved(editor, room):
    first = add_computer(editor, room["room"], hostname="pc-1")
    second = add_computer(editor, room["room"], hostname="pc-2")

    ok(editor.patch(f"/api/computers/{first}", json={"serial": "SN-1"}))
    ok(editor.patch(f"/api/computers/{second}", json={"serial": "SN-1"}))

    assert get_row(editor, second)["serial"] == "SN-1"


def test_bulk_update_new_fields(editor, room):
    ids = [
        add_computer(editor, room["room"], hostname="pc-1"),
        add_computer(editor, room["room"], hostname="pc-2"),
    ]

    for field, value in (("vnc", "x"), ("temp_until", "01.11.2026")):
        data = ok(editor.post("/api/computers/bulk-update", json={"ids": ids, "field": field, "value": value}))
        assert data["changed"] == ids

    for field in ("serial", "glpi_id"):
        response = editor.post("/api/computers/bulk-update", json={"ids": ids, "field": field, "value": "x"})
        assert response.status_code == 400, field


def test_export_has_new_columns(reader, editor, room):
    computer_id = add_computer(editor, room["room"], hostname="pc-1")
    ok(editor.patch(f"/api/computers/{computer_id}", json=NEW_FIELDS))

    response = reader.get("/api/export/computers.xlsx")
    assert response.status_code == 200
    ws = load_workbook(BytesIO(response.content))["Компьютеры"]

    headers = [cell.value for cell in ws[1]]
    assert headers[: len(COLUMNS)] == [column.short for column in COLUMNS]

    values = dict(zip(headers, [cell.value for cell in ws[2]]))
    assert values["Серийный"] == "SN-001"
    assert values["VNC"] == "tight"
    assert values["GLPI"] == "1234"
    assert values["Временно до"].date() == date(2026, 10, 15)  # в Excel – дата, не текст


def test_new_keys_reserved_for_user_fields(editor):
    for key in ("serial", "vnc", "glpi_id", "temp_until"):
        response = editor.post("/api/field-defs", json={"key": key, "label": "Поле"})
        assert response.status_code == 400, key


def test_temp_until_date_formats(editor, room):
    computer_id = add_computer(editor, room["room"], hostname="pc-1")
    year = date.today().year

    for text, expected in (
        ("15.10.2026", "15.10.2026"),
        ("1.2.27", "01.02.2027"),
        ("2026-12-31", "31.12.2026"),
        ("5/6", f"05.06.{year}"),
    ):
        data = ok(editor.patch(f"/api/computers/{computer_id}", json={"temp_until": text}))
        assert data["updated"]["temp_until"] == expected, text

    for text in ("до пятницы", "32.01.2026", "15.13.2026"):
        response = editor.patch(f"/api/computers/{computer_id}", json={"temp_until": text})
        assert response.status_code == 400, text
        assert "ДД.ММ.ГГГГ" in response.json()["detail"]

    history = ok(editor.get(f"/api/computers/{computer_id}"))["history"]
    assert history[0]["changes"]["temp_until"] == {"old": "31.12.2026", "new": f"05.06.{year}"}


def test_card_history_is_complete(editor, room):
    computer_id = add_computer(editor, room["room"], hostname="pc-1")

    for number in range(25):
        ok(editor.patch(f"/api/computers/{computer_id}", json={"note": f"n{number}"}))

    # 25 правок + «Создан»: карточка получает всю историю ПК, не 20 последних
    assert len(ok(editor.get(f"/api/computers/{computer_id}"))["history"]) == 26


def test_suggest_flag(reader):
    columns = ok(reader.get("/api/columns"))["columns"]
    suggest = {column["key"] for column in columns if column["suggest"]}

    # Подсказки – у столбцов с повторяющимися значениями в одну строку
    assert {"status", "type", "os", "model", "motherboard", "cpu", "ram", "gpu", "vnc", "gsit", "state", "label"} == suggest
