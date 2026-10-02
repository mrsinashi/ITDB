"""Сбор из GLPI / GSIT и сопоставление с ПК (этап 25). GLPI изображает
маленький HTTP-сервер в этом же процессе; данные ПК задаёт каждый тест."""
import json
import threading
import time
import urllib.parse
from datetime import datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
from cryptography.fernet import Fernet

import scan_normalize as sn
from conftest import add_computer, get_row, ok

KEY = Fernet.generate_key().decode()
NOW = datetime.now()


def ago(days):
    return (NOW - timedelta(days=days)).strftime("%Y-%m-%d %H:%M:%S")


@pytest.fixture(autouse=True)
def secret_key(monkeypatch):
    monkeypatch.setenv("ITDB_SECRET_KEY", KEY)


def pc(name, days=0, serial=None, uuid=None, ports=(), cpu="Intel(R) Core(TM) i5-10400 CPU @ 2.90GHz",
       memory=(8192,), disks=(("Samsung SSD 870 EVO 250GB", 250059),), gpus=(), soft=(), maker="HP",
       model="ProDesk 400 G7", os_name="Microsoft Windows 10 Pro"):
    return {
        "name": name, "date": ago(days) if days is not None else None, "serial": serial, "uuid": uuid,
        "ports": list(ports), "cpu": cpu, "memory": list(memory), "disks": list(disks), "gpus": list(gpus),
        "soft": list(soft), "maker": maker, "model": model, "os": os_name,
    }


class Glpi:
    computers = {}
    antivirus = []
    fresh_field = True
    fail_details = False
    disk_models = True   # отдавать ли список моделей дисков (DeviceHardDrive)


class FakeGlpi(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def send(self, status, body, headers=None):
        data = json.dumps(body).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        for name, value in (headers or {}).items():
            self.send_header(name, value)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        parts = urllib.parse.urlsplit(self.path)
        path = parts.path[len("/apirest.php/"):]
        query = urllib.parse.parse_qs(parts.query)

        if path == "initSession":
            return self.send(200, {"session_token": "s1"})

        if self.headers.get("Session-Token") != "s1":
            return self.send(401, ["ERROR_SESSION_TOKEN_INVALID", "нет сессии"])

        if path == "killSession":
            return self.send(200, {})

        if path == "getGlpiConfig":
            return self.send(200, {"cfg_glpi": {"version": "11.0.7"}})

        if path == "listSearchOptions/Computer":
            options = {
                "common": "Характеристики",
                "1": {"table": "glpi_computers", "field": "name"},
                "2": {"table": "glpi_computers", "field": "id"},
                "45": {"table": "glpi_operatingsystems", "field": "name"},
                "901": {"table": "glpi_agents", "field": "tag"},
            }
            if Glpi.fresh_field:
                options["9"] = {"table": "glpi_computers", "field": "last_inventory_update"}
            return self.send(200, options)

        if path == "search/Computer":
            start, end = (int(x) for x in query["range"][0].split("-"))
            rows = [
                {"1": c["name"], "2": cid, "9": c["date"], "45": c["os"], "901": "med"}
                for cid, c in sorted(Glpi.computers.items())
            ]
            page = rows[start:end + 1]
            return self.send(200, {"totalcount": len(rows), "count": len(page), "data": page})

        if path in ("ComputerAntivirus", "ItemAntivirus"):
            start, end = (int(x) for x in query["range"][0].split("-"))
            return self.send(200, Glpi.antivirus[start:end + 1])

        if path == "DeviceHardDrive":
            # модели дисков с интерфейсом: диск ПК — (название, объём[, интерфейс])
            if not Glpi.disk_models:
                return self.send(403, ["ERROR_RIGHT_MISSING", "нет права"])
            start, end = (int(x) for x in query["range"][0].split("-"))
            models = sorted({(d[0], d[2] if len(d) > 2 else 0) for c in Glpi.computers.values() for d in c["disks"]}, key=str)
            rows = [{"id": i + 1, "designation": n, "interfacetypes_id": kind} for i, (n, kind) in enumerate(models)]
            return self.send(200, rows[start:end + 1])

        if path.startswith("Computer/"):
            if Glpi.fail_details:
                return self.send(401, ["ERROR_SESSION_TOKEN_INVALID", "сессия кончилась"])
            cid = int(path.split("/")[1])
            c = Glpi.computers.get(cid)
            if c is None:
                return self.send(404, ["ERROR_ITEM_NOT_FOUND", "нет"])
            return self.send(200, {
                "id": cid, "name": c["name"], "serial": c["serial"], "uuid": c["uuid"],
                "manufacturers_id": c["maker"] or 0, "computermodels_id": c["model"] or 0,
                "_devices": {
                    "Item_DeviceProcessor": {"1": {"deviceprocessors_id": c["cpu"]}},
                    "Item_DeviceMemory": {str(i): {"size": m} for i, m in enumerate(c["memory"])},
                    "Item_DeviceHardDrive": {str(i): {"deviceharddrives_id": d[0], "capacity": d[1]} for i, d in enumerate(c["disks"])},
                    "Item_DeviceGraphicCard": {str(i): {"devicegraphiccards_id": g} for i, g in enumerate(c["gpus"])},
                },
                "_networkports": {
                    "NetworkPortEthernet": [
                        {"name": n, "mac": m, "NetworkName": {"IPAddress": [{"name": ip} for ip in ips]}}
                        for n, m, ips in c["ports"]
                    ],
                    "NetworkPortLocal": [{"name": "lo", "mac": "00:00:00:00:00:00"}],
                },
                "_softwares": [{"softwares_id": s} for s in c["soft"]],
            })

        return self.send(404, ["ERROR", "нет"])


@pytest.fixture(scope="module")
def glpi_url():
    server = ThreadingHTTPServer(("127.0.0.1", 0), FakeGlpi)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()


@pytest.fixture(autouse=True)
def reset_glpi():
    Glpi.computers = {}
    Glpi.antivirus = []
    Glpi.fresh_field = True
    Glpi.fail_details = False
    Glpi.disk_models = True


def setup_source(admin, url, kind="glpi", enabled=True):
    ok(admin.patch(f"/api/scan/sources/{kind}", json={"url": url, "login": "itdb", "password": "pw", "enabled": enabled}))


def collect(admin, kind="glpi"):
    run = ok(admin.post(f"/api/scan/sources/{kind}/collect"))
    for _ in range(200):
        run = ok(admin.get(f"/api/scan/runs/{run['id']}"))
        if run["status"] != "running":
            return run
        time.sleep(0.05)
    raise AssertionError("сбор не закончился")


def records(admin, kind="glpi"):
    data = ok(admin.get("/api/scan/records", params={"source": kind}))
    return data, {r["source_id"]: r for r in data["records"]}


ETH = "Realtek PCIe GbE Family Controller"


def add_pc(client, location_id, hostname, **fields):
    """ПК с любыми полями: POST принимает только расположение, имя и IP."""
    computer_id = add_computer(client, location_id, hostname=hostname)

    if fields:
        version = get_row(client, computer_id)["version"]
        ok(client.patch(f"/api/computers/{computer_id}", json=dict(fields, _version=version)))

    return computer_id


# ---------- Разбор значений ----------


def test_normalize_values():
    assert sn.clean_serial("Default string") is None
    assert sn.clean_serial("To be filled by O.E.M.") is None
    assert sn.clean_serial("1,90E+14") is None
    assert sn.clean_serial("0000000") is None
    assert sn.clean_serial(" 5CD1234ABC ") == "5CD1234ABC"
    assert sn.clean_uuid("03000200-0400-0500-0006-000700080009") is None
    assert sn.norm_mac("d8-bb-c1-00-00-01") == "D8:BB:C1:00:00:01"
    assert sn.is_virtual_mac("0A:00:27:00:00:05")      # VirtualBox host-only
    assert sn.is_virtual_mac("00:15:5D:01:02:03")      # Hyper-V
    assert not sn.is_virtual_mac("D8:BB:C1:00:00:01")
    assert sn.os_short("Microsoft Windows 10 Pro") == "Win 10"
    assert sn.os_short("Windows", "22H2", "10.0.22631") == "Win 11"
    assert sn.os_short("Windows", "22H2", "10.0.19045") == "Win 10"
    assert sn.os_short("Astra Linux CE") == "Astra"
    assert sn.cpu_short("Intel(R) Core(TM) i5-10400 CPU @ 2.90GHz") == "i5-10400"
    assert sn.cpu_short("Intel(R) Pentium(R) Gold G5400 CPU @ 3.70GHz") == "Pentium G5400"
    assert sn.cpu_short("AMD Ryzen 5 3400G with Radeon Vega Graphics") == "Ryzen 5 3400G"
    assert sn.ram_short(8192) == "8" and sn.ram_short(3968) == "4"
    assert sn.drives_short([
        {"name": "WDC WD10EZEX-08WN4A0", "mb": 1000204},
        {"name": "Samsung SSD 870 EVO 250GB", "mb": 250059},
        {"name": "Kingston DataTraveler USB Device", "mb": 15500},
    ]) == "SSD 250\nHDD 1TB"
    assert sn.gpu_short(["Intel(R) UHD Graphics 630", "NVIDIA GeForce GT 1030", "Radmin Mirror Driver V3"]) == "GT 1030"
    assert sn.model_short("Hewlett-Packard", "HP ProDesk 400 G6 MT") == "HP ProDesk 400 G6 MT"
    assert sn.model_short("System manufacturer", "System Product Name") is None
    assert sn.vnc_short(["TightVNC", "7-Zip"]) == "TightVNC"


# ---------- Сбор ----------


def test_collect_and_match(admin, editor, room, glpi_url):
    loc = room["room"]
    by_mac = add_pc(editor, loc, "ter-201-1", mac="D8:BB:C1:00:00:01")
    by_serial = add_pc(editor, loc, "hir-301-1", serial="MXL0000001")
    by_name = add_pc(editor, loc, "reg-101-1")
    wrong_mac = add_pc(editor, loc, "lab-101-1", mac="D8:BB:C1:99:99:99")
    by_id = add_pc(editor, loc, "adm-201-1", glpi_id="7")
    id_other_name = add_pc(editor, loc, "old-name", glpi_id="8")
    two_a = add_pc(editor, loc, "two-a", mac="D8:BB:C1:00:00:0A")
    two_b = add_pc(editor, loc, "two-b", serial="SER-TWO-B")
    archived = add_pc(editor, loc, "gone-pc", mac="D8:BB:C1:00:00:0C")
    ok(editor.post("/api/computers/archive", json={"ids": [archived], "archived": True}))

    Glpi.computers = {
        1: pc("ter-201-1", ports=[(ETH, "d8:bb:c1:00:00:01", ["10.0.2.11", "fe80::1"]),
                                  ("VirtualBox Host-Only Ethernet Adapter", "0a:00:27:00:00:05", ["192.168.56.1"]),
                                  ("TAP-Windows Adapter V9", "00:ff:12:34:56:78", ["10.8.0.6"])],
              soft=["TightVNC", "7-Zip"], gpus=["NVIDIA GeForce GT 1030"]),
        2: pc("hir-301-1", serial="MXL0000001"),
        3: pc("reg-101-1", serial="Default string", ports=[(ETH, "d8:bb:c1:00:00:03", [])]),
        4: pc("lab-101-1", ports=[(ETH, "d8:bb:c1:00:00:04", [])]),
        7: pc("adm-201-1"),
        8: pc("new-name"),
        9: pc("two", serial="SER-TWO-B", ports=[(ETH, "d8:bb:c1:00:00:0a", [])]),
        10: pc("gone-pc", ports=[(ETH, "d8:bb:c1:00:00:0c", [])]),
        11: pc("unknown-pc", ports=[(ETH, "d8:bb:c1:00:00:0b", [])]),
        # Устаревшая и без даты — не берутся
        20: pc("stale-pc", days=10, ports=[(ETH, "d8:bb:c1:00:00:20", [])]),
        21: pc("manual-pc", days=None),
        # Дубль: то же железо (MAC), запись старее — остаётся 23
        22: pc("dup-pc", days=2, ports=[(ETH, "d8:bb:c1:00:00:22", [])]),
        23: pc("dup-pc", days=0, ports=[(ETH, "d8:bb:c1:00:00:22", [])]),
        # Серийный на трёх записях — общий, не признак (и не склеивает)
        30: pc("s-1", serial="US00040148", ports=[(ETH, "d8:bb:c1:00:00:31", [])]),
        31: pc("s-2", serial="US00040148", ports=[(ETH, "d8:bb:c1:00:00:32", [])]),
        32: pc("s-3", serial="US00040148", ports=[(ETH, "d8:bb:c1:00:00:33", [])]),
    }
    Glpi.antivirus = [
        {"computers_id": 1, "name": "Kaspersky Endpoint Security", "is_active": 1, "is_uptodate": 1, "antivirus_version": "12"},
        {"computers_id": 1, "name": "Windows Defender", "is_active": 0, "is_uptodate": 1},
        {"computers_id": 2, "name": "Old AV", "is_active": 1, "is_uptodate": 0, "is_deleted": 1},
    ]
    setup_source(admin, glpi_url)
    run = collect(admin)

    assert run["status"] == "ok", run["message"]
    stats = run["stats"]
    assert (stats["total"], stats["fresh"], stats["stale"], stats["no_date"]) == (16, 14, 1, 1)
    assert stats["shared"] == 1 and stats["dup"] == 1
    assert "progress" not in stats

    data, recs = records(admin)
    assert data["web_url"] == glpi_url
    assert set(recs) == {1, 2, 3, 4, 7, 8, 9, 10, 11, 22, 23, 30, 31, 32}

    r = recs[1]
    assert (r["state"], r["computer_id"], r["by"]) == ("key", by_mac, ["mac"])
    assert r["values"]["ip"] == "10.0.2.11"                  # без IPv6, VirtualBox, VPN
    assert r["values"]["mac"] == "D8:BB:C1:00:00:01"
    assert r["values"]["os"] == "Win 10" and r["values"]["cpu"] == "i5-10400" and r["values"]["ram"] == "8"
    assert r["values"]["drive"] == "SSD 250" and r["values"]["gpu"] == "GT 1030" and r["values"]["vnc"] == "TightVNC"
    assert r["values"]["model"] == "HP ProDesk 400 G7"
    assert [a["name"] for a in r["antivirus"]] == ["Kaspersky Endpoint Security", "Windows Defender"]
    assert r["antivirus"][1]["active"] is False
    assert recs[2]["antivirus"] == []                          # удалённый антивирус не в счёт

    assert (recs[2]["state"], recs[2]["computer_id"], recs[2]["by"]) == ("key", by_serial, ["serial"])
    assert recs[3]["values"]["serial"] is None                 # «Default string»
    assert (recs[3]["state"], recs[3]["computer_id"]) == ("name", by_name)
    assert recs[4]["state"] == "conflict" and "MAC" in recs[4]["note"]
    assert recs[4]["candidates"] == [wrong_mac]
    assert (recs[7]["state"], recs[7]["computer_id"], recs[7]["by"]) == ("key", by_id, ["id"])
    assert (recs[8]["state"], recs[8]["computer_id"]) == ("name", id_other_name)   # только GLPI ID
    assert recs[9]["state"] == "conflict" and set(recs[9]["candidates"]) == {two_a, two_b}
    assert recs[10]["state"] == "none" and "gone-pc" in recs[10]["note"]            # ПК в архиве
    assert recs[11]["state"] == "none" and recs[11]["note"] is None
    assert recs[22]["state"] == "dup" and recs[22]["dup_of"] == 23
    assert recs[23]["state"] == "none"
    assert all(recs[i]["values"]["serial"] == "US00040148" for i in (30, 31, 32))
    assert all(recs[i]["state"] == "none" for i in (30, 31, 32))                    # серийный не признак
    assert data["counts"] == {"key": 3, "link": 0, "name": 2, "conflict": 2, "none": 6, "dup": 1}

    # Ничего не записано в ПК (правило 1)
    row = get_row(editor, by_mac)
    assert not row["ip"] and not row["os"] and not row["cpu"]


def test_match_follows_table_edits(admin, editor, room, glpi_url):
    """Сопоставление считается по текущим данным таблицы, без пересбора."""
    computer = add_pc(editor, room["room"], "other-name")
    Glpi.computers = {1: pc("pc-1", ports=[(ETH, "d8:bb:c1:00:00:01", [])])}
    setup_source(admin, glpi_url)
    collect(admin)
    assert records(admin)[1][1]["state"] == "none"

    version = get_row(editor, computer)["version"]
    ok(editor.patch(f"/api/computers/{computer}", json={"mac": "d8:bb:c1:00:00:01", "_version": version}))
    r = records(admin)[1][1]
    assert (r["state"], r["computer_id"]) == ("key", computer)


def test_two_records_same_name(admin, editor, room, glpi_url):
    """Два ПК с одним именем в GLPI: один сопоставлен по MAC, второй — конфликт."""
    computer = add_pc(editor, room["room"], "glaz-kdp3-3", mac="D8:BB:C1:00:00:04")
    Glpi.computers = {
        4: pc("glaz-kdp3-3", ports=[(ETH, "d8:bb:c1:00:00:04", [])]),
        5: pc("glaz-kdp3-3", serial="PF2ABCDE", ports=[(ETH, "d8:bb:c1:00:00:05", [])]),
    }
    setup_source(admin, glpi_url)
    collect(admin)
    recs = records(admin)[1]
    assert (recs[4]["state"], recs[4]["computer_id"]) == ("key", computer)
    assert recs[5]["state"] == "conflict" and recs[5]["candidates"] == [computer]


def test_decisions(admin, editor, reader, room, glpi_url):
    first = add_pc(editor, room["room"], "pc-1")
    second = add_pc(editor, room["room"], "pc-2", mac="D8:BB:C1:00:00:09")
    Glpi.computers = {1: pc("pc-1", ports=[(ETH, "d8:bb:c1:00:00:01", [])])}
    setup_source(admin, glpi_url)
    collect(admin)
    assert records(admin)[1][1]["state"] == "name"

    # Только администратор
    assert editor.post("/api/scan/records/glpi/1/link", json={"computer_id": first}).status_code == 403
    assert reader.get("/api/scan/records", params={"source": "glpi"}).status_code == 403
    assert editor.get("/api/scan/records", params={"source": "glpi"}).status_code == 200   # «Проверка» (26г)

    # «Не этот ПК» — больше не предлагается
    ok(admin.post("/api/scan/records/glpi/1/reject", json={"computer_id": first}))
    r = records(admin)[1][1]
    assert r["state"] == "none"
    assert [(d["action"], d["hostname"]) for d in r["decisions"]] == [("reject", "pc-1")]

    # «Это этот ПК» — главнее признаков (у pc-2 MAC другой)
    ok(admin.post("/api/scan/records/glpi/1/link", json={"computer_id": second}))
    r = records(admin)[1][1]
    assert (r["state"], r["computer_id"]) == ("link", second)

    # Переживает повторный сбор
    collect(admin)
    assert records(admin)[1][1]["state"] == "link"

    # Забыть решения — снова по признакам (имени)
    ok(admin.delete("/api/scan/records/glpi/1/decisions"))
    r = records(admin)[1][1]
    assert (r["state"], r["computer_id"], r["decisions"]) == ("name", first, [])

    # В архив привязать нельзя; нет записи — 404
    ok(editor.post("/api/computers/archive", json={"ids": [second], "archived": True}))
    assert admin.post("/api/scan/records/glpi/1/link", json={"computer_id": second}).status_code == 400
    assert admin.post("/api/scan/records/glpi/99/link", json={"computer_id": first}).status_code == 404

    # История — у администратора, без значений ПК
    items = ok(admin.get("/api/history", params={"entity": "scan_records"}))["items"]
    fields = [list(item["changes"]) for item in items]
    assert sorted(f[0] for f in fields) == ["link", "reject", "reset"]
    assert all(item["title"] == "GLPI №1 pc-1" for item in items)
    assert ok(editor.get("/api/history"))["total"] == len(ok(editor.get("/api/history"))["items"])
    assert not [i for i in ok(editor.get("/api/history"))["items"] if i["entity"] == "scan_records"]


def test_collect_refused(admin, editor, glpi_url):
    assert admin.post("/api/scan/sources/glpi/collect").status_code == 400          # нет адреса
    setup_source(admin, glpi_url, enabled=False)
    response = admin.post("/api/scan/sources/glpi/collect")
    assert response.status_code == 400 and "выключен" in response.json()["detail"]
    assert admin.post("/api/scan/sources/jabber/collect").status_code == 400        # Jabber — тоже (этап 26д)
    assert admin.post("/api/scan/sources/nope/collect").status_code == 404
    assert admin.get("/api/scan/records", params={"source": "jabber"}).status_code == 404
    assert editor.post("/api/scan/sources/glpi/collect").status_code == 403


def test_collect_errors_keep_old_records(admin, glpi_url):
    Glpi.computers = {1: pc("pc-1", ports=[(ETH, "d8:bb:c1:00:00:01", [])])}
    setup_source(admin, glpi_url)
    assert collect(admin)["status"] == "ok"

    # Сбой на середине — прежние записи остаются
    Glpi.fail_details = True
    run = collect(admin)
    assert run["status"] == "error" and run["message"]
    assert list(records(admin)[1]) == [1]

    # Нет поля даты проверки — сбор не начинается
    Glpi.fail_details = False
    Glpi.fresh_field = False
    run = collect(admin)
    assert run["status"] == "error" and "дат" in run["message"]

    # Последний запуск — в настройках источника, журнал — новые сверху
    source = next(s for s in ok(admin.get("/api/scan/sources"))["sources"] if s["kind"] == "glpi")
    assert source["last_run"]["id"] == run["id"]
    runs = ok(admin.get("/api/scan/runs", params={"source": "glpi"}))
    assert [r["status"] for r in runs] == ["error", "error", "ok"]


# ---------- Названия (этап 25б) ----------

import scan_values as sv  # noqa: E402


def test_same_by_meaning():
    assert sv.same_line("os", "Win 10", "Windows 10")
    assert not sv.same_line("os", "Win 10", "Windows 11")
    assert sv.same_line("os", "Майкрософт Windows 10 Pro", "Windows 10")
    assert sv.same_line("cpu", "Pentium J3710", "Intel Pentium J3710")
    assert not sv.same_line("cpu", "i5-10400", "i5-10500")
    assert sv.same_line("model", "Lenovo ThinkCentre M600", "M600")
    assert not sv.same_line("model", "HP ProDesk 400 G7", "HP ProDesk 400 G6")
    assert sv.same_line("vnc", "TightVNC", "Tight")
    assert not sv.same_line("vnc", "UltraVNC", "Tight")
    assert not sv.same_line("drive", "SSD 256", "SSD 250")    # 26е: объём точный — разные диски
    assert sv.same_line("drive", "HDD 1TB", "HDD 1000")
    assert not sv.same_line("drive", "SSD 240", "SSD 250")    # разные диски
    assert sv.same_line("drive", "HDD 500", "500")
    assert sv.same_line("ram", "4", "4 ГБ")

    names = sv.Names()
    mac = names.compare("mac", "6C:4B:90:58:66:26\n00:19:D2:F8:32:89", "6c:4b:90:58:66:26")
    assert mac["mark"] == "≈"                                  # в таблице записана часть
    os_ = names.compare("os", "Win 10", "Windows 10")
    assert (os_["source"], os_["raw"], os_["mark"]) == ("Windows 10", "Win 10", "=")
    assert names.compare("hostname", "glaz-kdp3-3", "glaz-kdp3-1")["mark"] == "≠"
    assert names.compare("vnc", "TightVNC", "")["mark"] == ""


def test_names_learned_and_manual():
    pairs = [
        ("cpu", "Pentium J3710", "Intel Pentium J3710"),
        ("cpu", "Pentium J3710", "Intel Pentium J3710"),
        ("drive", "HDD 500", "HDD 500"),
        ("drive", "HDD 500", "HDD 500"),
        ("drive", "HDD 500", "500"),              # небрежно у одного — не учится
    ]
    names = sv.Names(
        aliases=[{"field": "os", "source": "Win 10", "table": "Десятка", "kind": "same"},
                 {"field": "model", "source": "HP 400", "table": "HP", "kind": "differ"}],
        table_values={"model": {"M600": 5}},
        choices={"gpu": ["GeForce GT 1030"]},
        pairs=pairs,
    )
    assert names.canonical_line("cpu", "Pentium J3710") == ("Intel Pentium J3710", "learned")
    assert names.canonical_line("drive", "HDD 500") == ("HDD 500", "")
    assert names.canonical_line("os", "Win 10") == ("Десятка", "manual")
    assert names.compare("os", "Win 10", "Десятка")["mark"] == "="
    assert names.canonical_line("model", "Lenovo ThinkCentre M600") == ("M600", "table")
    assert names.canonical_line("gpu", "GT 1030") == ("GeForce GT 1030", "table")
    assert names.compare("model", "HP 400", "HP")["mark"] == "≠"          # «это разное»


def test_names_api(admin, editor, reader, room, glpi_url):
    add_pc(editor, room["room"], "pc-1", mac="D8:BB:C1:00:00:01", os="Windows 10", cpu="Intel Pentium J3710", model="M600")
    Glpi.computers = {
        1: pc("pc-1", ports=[(ETH, "d8:bb:c1:00:00:01", [])], cpu="Intel(R) Pentium(R) CPU J3710 @ 1.60GHz",
              maker="LENOVO", model="ThinkCentre M600"),
        2: pc("pc-2", ports=[(ETH, "d8:bb:c1:00:00:02", [])], cpu="Intel(R) Pentium(R) CPU J3710 @ 1.60GHz"),
    }
    setup_source(admin, glpi_url)
    collect(admin)
    recs = records(admin)[1]
    rows = {c["field"]: c for c in recs[1]["compare"]}
    assert (rows["os"]["source"], rows["os"]["raw"], rows["os"]["mark"]) == ("Windows 10", "Win 10", "=")
    assert (rows["cpu"]["source"], rows["cpu"]["mark"]) == ("Intel Pentium J3710", "=")
    assert (rows["model"]["source"], rows["model"]["mark"]) == ("M600", "=")
    # У другого ПК (нет в ITDB) — те же названия, как в таблице
    other = {c["field"]: c for c in recs[2]["compare"]}
    assert (other["cpu"]["source"], other["cpu"]["how"]) == ("Intel Pentium J3710", "learned")
    assert other["os"]["source"] == "Windows 10"

    learned = ok(admin.get("/api/scan/names"))
    assert {(n["field"], n["source"], n["table"], n["origin"]) for n in learned} >= {
        ("os", "Win 10", "Windows 10", "learned"), ("cpu", "Pentium J3710", "Intel Pentium J3710", "learned"),
    }

    # Вручную: «одно и то же» — главнее, «это разное» — не считать одинаковыми
    ok(admin.post("/api/scan/names", json={"field": "os", "source": "Win 10", "table": "Win10 Pro"}))
    ok(admin.post("/api/scan/names", json={"field": "cpu", "source": "Pentium J3710", "table": "Intel Pentium J3710", "kind": "differ"}))
    recs = records(admin)[1]
    rows = {c["field"]: c for c in recs[1]["compare"]}
    assert rows["cpu"]["mark"] == "≠"
    assert {c["field"]: c for c in recs[2]["compare"]}["os"]["source"] == "Win10 Pro"

    manual = [n for n in ok(admin.get("/api/scan/names")) if n["origin"] == "manual"]
    assert len(manual) == 2
    assert admin.post("/api/scan/names", json={"field": "hostname", "source": "a", "table": "b"}).status_code == 400
    # Редактор тоже решает (режим «Расхождения»), только чтение — нет
    assert reader.post("/api/scan/names", json={"field": "os", "source": "a", "table": "b"}).status_code == 403
    for n in manual:
        ok(admin.delete(f"/api/scan/names/{n['id']}"))
    assert not [n for n in ok(admin.get("/api/scan/names")) if n["origin"] == "manual"]
    items = ok(admin.get("/api/history", params={"entity": "scan_aliases"}))["items"]
    assert len(items) == 4 and items[0]["title"].startswith(("OS: ", "CPU: "))


# ---------- Расхождения (этап 26) ----------


def diffs_of(client, rejected=False):
    return ok(client.get("/api/scan/diffs", params={"rejected": rejected}))


def test_keep_is_not_same():
    """«В таблице своё» (ЕГИСЗ): пара — не расхождение, но ничему не учит."""
    names = sv.Names(
        aliases=[{"field": "model", "source": "ASUS PRIME H370-PLUS", "table": "ЕГИСЗ", "kind": "keep"}],
        table_values={"model": {"ЕГИСЗ": 30}},
        pairs=[("model", "ASUS PRIME H370-PLUS", "ЕГИСЗ")] * 30,
    )
    row = names.compare("model", "ASUS PRIME H370-PLUS", "ЕГИСЗ")
    assert (row["mark"], row["manual"], row["source"]) == ("=", "keep", "ЕГИСЗ")
    # ПК без значения в таблице — получит значение источника, не «ЕГИСЗ»
    assert names.canonical_line("model", "ASUS PRIME H370-PLUS") == ("ASUS PRIME H370-PLUS", "")
    # Другое значение таблицы при той же плате — расхождение
    assert names.compare("model", "ASUS PRIME H370-PLUS", "Aquarius")["mark"] == "≠"


def test_diffs_accept_reject_keep(admin, editor, reader, room, glpi_url):
    loc = room["room"]
    first = add_pc(editor, loc, "egisz-1", mac="04:D9:F5:00:00:01", model="ЕГИСЗ", ram="16")
    second = add_pc(editor, loc, "egisz-2", mac="04:D9:F5:00:00:02", model="ЕГИСЗ")
    Glpi.computers = {
        1: pc("egisz-1", ports=[(ETH, "04:d9:f5:00:00:01", ["10.0.1.60"])], maker="ASUSTeK COMPUTER INC.", model="PRIME H370-PLUS"),
        2: pc("egisz-2", ports=[(ETH, "04:d9:f5:00:00:02", [])], maker="ASUSTeK COMPUTER INC.", model="PRIME H370-PLUS"),
    }
    setup_source(admin, glpi_url)
    collect(admin)

    # Смотреть пометки может и «только чтение», решать — нет
    assert reader.get("/api/scan/diffs").status_code == 200
    assert reader.post("/api/scan/diffs/reject", json={"items": []}).status_code == 403
    data = diffs_of(editor)
    items = {(d["computer_id"], d["field"]): d for d in data["items"]}
    assert data["computers"] == 2 and data["sources"] == ["glpi"]
    model = items[(first, "model")]
    assert (model["table"], model["proposed"], model["kind"], model["same_pair"]) == ("ЕГИСЗ", "ASUS PRIME H370-PLUS", "diff", 1)
    assert items[(first, "ip")]["kind"] == "fill" and items[(first, "ram")]["proposed"] == "8"
    assert items[(first, "os")]["proposed"] == "Win 10"
    # Признаки доверия: как сопоставлен и что предлагает источник
    src = model["sources"][0]
    assert (src["title"], src["state"], src["by"], src["value"]) == ("GLPI", "key", ["mac"], "ASUS PRIME H370-PLUS")

    # «В таблице своё» — у обоих ПК модель больше не расхождение
    ok(editor.post("/api/scan/names", json={"field": "model", "source": "ASUS PRIME H370-PLUS", "table": "ЕГИСЗ", "kind": "keep"}))
    assert not [d for d in diffs_of(editor)["items"] if d["field"] == "model"]

    # Отклонить RAM у первого — скрыто, но видно среди отклонённых; вернуть
    ok(editor.post("/api/scan/diffs/reject", json={"items": [{"computer_id": first, "field": "ram", "raw": "8"}]}))
    data = diffs_of(editor)
    assert (first, "ram") not in {(d["computer_id"], d["field"]) for d in data["items"]} and data["rejected"] == 1
    shown = [d for d in diffs_of(editor, True)["items"] if d["rejected_by"]]
    assert [(d["field"], d["rejected_by"]) for d in shown] == [("ram", "editor")]
    ok(editor.post("/api/scan/diffs/unreject", json={"items": [{"computer_id": first, "field": "ram", "raw": "8"}]}))
    assert (first, "ram") in {(d["computer_id"], d["field"]) for d in diffs_of(editor)["items"]}

    # Принять IP и OS — в ПК и в Историю с пометкой источника
    result = ok(editor.post("/api/scan/diffs/accept", json={"items": [
        {"computer_id": first, "field": "ip", "value": "10.0.1.60", "table": "", "source": "GLPI №1"},
        {"computer_id": first, "field": "os", "value": "Win 10", "table": "", "source": "GLPI №1"},
        {"computer_id": second, "field": "os", "value": "Win 10", "table": "уже другое", "source": "GLPI №2"},
    ]}))
    assert result["accepted"] == 2 and len(result["skipped"]) == 1
    row = get_row(editor, first)
    assert (row["ip"], row["os"]) == ("10.0.1.60", "Win 10")
    history = ok(editor.get("/api/history", params={"entity": "computers", "entity_id": first}))["items"]
    assert history[0]["changes"]["ip"]["scan"] == "GLPI №1"
    assert {(d["computer_id"], d["field"]) for d in diffs_of(editor)["items"]} >= {(second, "os")}
    assert (first, "os") not in {(d["computer_id"], d["field"]) for d in diffs_of(editor)["items"]}

    # Выключенный источник — расхождений нет
    ok(admin.patch("/api/scan/sources/glpi", json={"enabled": False}))
    assert diffs_of(editor)["items"] == []


def test_names_cycle_auto(admin, editor, room, glpi_url):
    """Двойной клик по отметке: «одно и то же» → «в таблице своё» → как решит сравнение."""
    add_pc(editor, room["room"], "pc-1", mac="D8:BB:C1:00:00:01", model="ЕГИСЗ")
    Glpi.computers = {1: pc("pc-1", ports=[(ETH, "d8:bb:c1:00:00:01", [])], maker="ASUS", model="PRIME H370-PLUS")}
    setup_source(admin, glpi_url)
    collect(admin)

    def model_row():
        return {c["field"]: c for c in records(admin)[1][1]["compare"]}["model"]

    row = model_row()
    assert (row["mark"], row["manual"], row["auto_equal"]) == ("≠", None, False)
    body = {"field": "model", "source": row["raw"], "table": "ЕГИСЗ"}
    ok(admin.post("/api/scan/names", json=dict(body, kind="same")))
    assert (model_row()["mark"], model_row()["manual"]) == ("=", "same")
    ok(admin.post("/api/scan/names", json=dict(body, kind="keep")))
    assert (model_row()["mark"], model_row()["manual"]) == ("=", "keep")
    assert len([n for n in ok(admin.get("/api/scan/names")) if n["origin"] == "manual"]) == 1
    ok(admin.post("/api/scan/names", json=dict(body, kind="auto")))
    assert (model_row()["mark"], model_row()["manual"]) == ("≠", None)
    assert not [n for n in ok(admin.get("/api/scan/names")) if n["origin"] == "manual"]


def test_canonical_keeps_model_numbers():
    """Название из таблицы подставляется, только если в нём все «номерные» части."""
    names = sv.Names(table_values={"model": {"HP ProDesk 400": 5, "M600": 3}})
    assert names.canonical_line("model", "HP ProDesk 400 G6 MT") == ("HP ProDesk 400 G6 MT", "")
    assert names.canonical_line("model", "Lenovo ThinkCentre M600") == ("M600", "table")
    # У самого ПК менее подробное значение — не расхождение
    assert names.compare("model", "HP ProDesk 400 G6 MT", "HP ProDesk 400")["mark"] == "="



def test_diffs_unsure_and_partial(admin, editor, room, glpi_url):
    """VNC: просмотрщик не сервер, TightVNC рядом с другим — TightVNC; два других —
    неточно. В таблице часть (≈) — не в счёт; ручное «=» у многострочного значения."""
    assert sn.vnc_short(["TightVNC", "UltraVNC 1.4"]) == "TightVNC"     # решение пользователя 01.10
    assert sn.vnc_short(["UltraVNC 1.4", "RealVNC Server"]) == "UltraVNC\nRealVNC"

    loc = room["room"]
    pc1 = add_pc(editor, loc, "vnc-1", mac="04:D9:F5:00:01:01")
    pc2 = add_pc(editor, loc, "vnc-2", mac="04:D9:F5:00:01:02", vnc="UltraVNC")
    pc3 = add_pc(editor, loc, "vnc-3", mac="04:D9:F5:00:01:03", vnc="Tight")
    Glpi.computers = {
        1: pc("vnc-1", ports=[(ETH, "04:d9:f5:00:01:01", [])], soft=["UltraVNC 1.4", "RealVNC Server"]),
        2: pc("vnc-2", ports=[(ETH, "04:d9:f5:00:01:02", [])], soft=["UltraVNC 1.4", "RealVNC Server"]),
        3: pc("vnc-3", ports=[(ETH, "04:d9:f5:00:01:03", [])], soft=["TightVNC", "UltraVNC Viewer 1.4", "UltraVNC 1.4"]),
    }
    setup_source(admin, glpi_url)
    collect(admin)
    assert records(admin)[1][3]["values"]["vnc"] == "TightVNC"

    data = diffs_of(editor)
    items = {(d["computer_id"], d["field"]): d for d in data["items"]}
    vnc1 = items[(pc1, "vnc")]
    assert vnc1["kind"] == "unsure" and "VNC" in vnc1["unsure"]
    assert items[(pc2, "vnc")]["kind"] == "partial"                     # «UltraVNC» — один из двух
    assert (pc3, "vnc") not in items                                    # Tight = TightVNC
    counted = {d["computer_id"] for d in data["items"] if d["kind"] != "partial" and not d.get("rejected_by")}
    assert data["computers"] == len(counted)                            # partial не в счёт ПК

    # Двойной клик по «≈» у многострочного: «одно и то же» целиком — отметка «=»
    rec = records(admin)[1][2]
    vnc = next(c for c in rec["compare"] if c["field"] == "vnc")
    assert vnc["mark"] == "≈" and vnc["manual"] is None and vnc["auto_equal"] is False
    ok(editor.post("/api/scan/names", json={"field": "vnc", "source": vnc["raw"], "table": "UltraVNC", "kind": "same"}))
    rec = records(admin)[1][2]
    vnc = next(c for c in rec["compare"] if c["field"] == "vnc")
    assert vnc["mark"] == "=" and vnc["manual"] == "same"
    assert (pc2, "vnc") not in {(d["computer_id"], d["field"]) for d in diffs_of(editor)["items"]}


def test_marks_settings(admin, editor, reader):
    marks = ok(reader.get("/api/scan/marks"))
    assert [m["kind"] for m in marks] == ["diff", "fill", "unsure", "partial"]
    assert marks[0]["label"] == "Отличается" and marks[3]["enabled"] is False

    assert reader.patch("/api/scan/marks/diff", json={"bold": True}).status_code == 403
    assert editor.patch("/api/scan/marks/diff", json={"color": "red"}).status_code == 400
    saved = ok(editor.patch("/api/scan/marks/diff", json={"strike": True, "frame": "#cc0000", "always": True}))
    assert (saved["strike"], saved["frame"], saved["always"]) == (True, "#cc0000", True)
    items = ok(admin.get("/api/history", params={"entity": "scan_marks"}))["items"]
    assert items[0]["title"] == "Отличается" and set(items[0]["changes"]) == {"strike", "frame", "always"}

    # Вернуть как было — таблица пометок между тестами не очищается
    ok(editor.patch("/api/scan/marks/diff", json={"strike": False, "frame": "", "always": False}))


# ---------- Столбец «Антивирусы» (этап 26д) ----------


def test_antivirus_column(admin, editor, reader, room, glpi_url):
    from io import BytesIO

    from openpyxl import load_workbook

    from api_scan_diffs import pick_antivirus

    loc = room["room"]
    first = add_pc(editor, loc, "pc-1", mac="D8:BB:C1:00:00:01")
    second = add_pc(editor, loc, "pc-2", mac="D8:BB:C1:00:00:02")
    Glpi.computers = {
        1: pc("pc-1", ports=[(ETH, "d8:bb:c1:00:00:01", [])]),
        2: pc("pc-2", ports=[(ETH, "d8:bb:c1:00:00:02", [])]),
    }
    Glpi.antivirus = [
        {"id": 1, "computers_id": 1, "name": "Kaspersky Endpoint Security", "is_active": 1, "is_uptodate": 1, "antivirus_version": "11.11"},
        {"id": 2, "computers_id": 1, "name": "Windows Defender", "is_active": 0, "is_uptodate": 1},
        {"id": 3, "computers_id": 2, "name": "Kaspersky Endpoint Security", "is_active": 1, "is_uptodate": 0},
    ]
    setup_source(admin, glpi_url)
    assert collect(admin)["status"] == "ok"

    # Антивирусы ПК — вместе с расхождениями (видят все), состояние словом
    av = ok(reader.get("/api/scan/diffs"))["antivirus"]
    assert [(a["name"], a["status"]) for a in av[str(first)]] == [("Kaspersky Endpoint Security", "on"), ("Windows Defender", "off")]
    assert av[str(first)][0]["version"] == "11.11" and av[str(first)][0]["source"] == "GLPI №1"
    assert [a["status"] for a in av[str(second)]] == ["old"]

    # Главнее GLPI; нет у него антивирусов — берутся из GSIT
    x, y = {"name": "x"}, {"name": "y"}
    assert pick_antivirus({5: {"glpi": [y], "gsit": [x]}}) == {5: [y]}
    assert pick_antivirus({5: {"glpi": [], "gsit": [x]}}) == {5: [x]}

    # В строке ПК столбца нет, править его нельзя
    row = get_row(editor, first)
    assert "antivirus" not in row
    assert editor.patch(f"/api/computers/{first}", json={"antivirus": "x", "_version": row["version"]}).status_code == 400

    # Вид столбца: по умолчанию показываются все состояния, скрытых нет
    settings = ok(reader.get("/api/scan/antivirus"))
    assert [s["kind"] for s in settings["statuses"]] == ["on", "old", "off"]
    assert all(s["show"] for s in settings["statuses"]) and settings["hidden"] == []
    assert reader.patch("/api/scan/antivirus", json={"kind": "on", "bold": True}).status_code == 403
    assert editor.patch("/api/scan/antivirus", json={"kind": "on", "color": "green"}).status_code == 400
    assert editor.patch("/api/scan/antivirus", json={"kind": "nope", "bold": True}).status_code == 400
    settings = ok(editor.patch("/api/scan/antivirus", json={"kind": "old", "show": False, "bold": True}))
    old = next(s for s in settings["statuses"] if s["kind"] == "old")
    assert (old["show"], old["bold"]) == (False, True)
    settings = ok(editor.patch("/api/scan/antivirus", json={"hidden": ["Windows Defender", "windows defender", " "]}))
    assert settings["hidden"] == ["Windows Defender"]
    assert ok(reader.get("/api/scan/antivirus"))["hidden"] == ["Windows Defender"]
    items = ok(reader.get("/api/history", params={"entity": "scan_antivirus"}))["items"]
    assert {i["title"] for i in items} == {"Базы устарели", "Не показывать"}

    # Выгрузка: показываемые антивирусы строками (Defender скрыт, «базы устарели» не показываются)
    response = reader.get("/api/export/computers.xlsx")
    ws = load_workbook(BytesIO(response.content))["Компьютеры"]
    headers = [cell.value for cell in ws[1]]
    values = {row[headers.index("HOSTNAME")].value: row[headers.index("Антивирус")].value for row in ws.iter_rows(min_row=2)}
    assert values == {"pc-1": "Kaspersky Endpoint Security", "pc-2": None}

    ok(editor.patch("/api/scan/antivirus", json={"kind": "old", "show": True}))
    ok(editor.patch("/api/scan/antivirus", json={"hidden": []}))
    response = reader.get("/api/export/computers.xlsx")
    ws = load_workbook(BytesIO(response.content))["Компьютеры"]
    values = {row[headers.index("HOSTNAME")].value: row[headers.index("Антивирус")].value for row in ws.iter_rows(min_row=2)}
    assert values == {
        "pc-1": "Kaspersky Endpoint Security\nWindows Defender — выключен",
        "pc-2": "Kaspersky Endpoint Security — базы устарели",
    }
