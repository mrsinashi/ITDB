"""Выгрузка в Excel: доступна любой роли, архив – отдельным листом."""
from io import BytesIO

from openpyxl import load_workbook

from conftest import add_computer, ok


def workbook(response):
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content))


def column_values(ws, header):
    headers = [cell.value for cell in ws[1]]
    index = headers.index(header)
    return [row[index].value for row in ws.iter_rows(min_row=2)]


def test_export_without_archive(reader, editor, room):
    add_computer(editor, room["room"], hostname="a")
    archived_id = add_computer(editor, room["room"], hostname="b")
    ok(editor.post("/api/computers/archive", json={"ids": [archived_id], "archived": True}))

    wb = workbook(reader.get("/api/export/computers.xlsx"))
    assert wb.sheetnames == ["Компьютеры"]
    assert column_values(wb["Компьютеры"], "HOSTNAME") == ["a"]


def test_export_with_archive(reader, editor, room):
    add_computer(editor, room["room"], hostname="a")
    archived_id = add_computer(editor, room["room"], hostname="b")
    ok(editor.post("/api/computers/archive", json={"ids": [archived_id], "archived": True}))

    response = reader.get("/api/export/computers.xlsx", params={"archive": "true"})
    assert "archive" in response.headers["content-disposition"]

    wb = workbook(response)
    assert wb.sheetnames == ["Компьютеры", "Архив"]
    assert column_values(wb["Компьютеры"], "HOSTNAME") == ["a"]
    assert column_values(wb["Архив"], "HOSTNAME") == ["b"]


def test_export_has_user_fields(editor, room):
    ok(editor.post("/api/field-defs", json={"key": "phone", "label": "Телефон"}))
    computer_id = add_computer(editor, room["room"], hostname="a")
    ok(editor.patch(f"/api/computers/{computer_id}", json={"phone": "12-34"}))

    ws = workbook(editor.get("/api/export/computers.xlsx"))["Компьютеры"]
    headers = [cell.value for cell in ws[1]]
    # Пользовательское поле – перед «Статус», как в таблице
    assert headers.index("Телефон") == headers.index("Статус") - 1
    assert column_values(ws, "Телефон") == ["12-34"]
