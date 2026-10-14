"""Этап 57: записи GLPI / GSIT не стираются (сбой, урезанный ответ, удалённый в источнике
ПК), давние ничего не предлагают и уступают действующим; убрать номер GSIT, когда на ПК
работает GLPI, а GSIT молчит; убранный номер GSIT больше не предлагается."""
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from conftest import add_pc, diffs_data, diffs_of, get_row, history, ok
from db import SessionLocal
from fakes import ETH, Glpi, collect, pc, setup_source
from models import ScanJabberUser, ScanRecord
from scan_match import match_all

MAC = ["d8:bb:c1:00:00:0" + str(i) for i in range(10)]


def port(i, *ips):
    return [(ETH, MAC[i], list(ips))]


def records(client, kind="glpi"):
    return {r["source_id"]: r for r in ok(client.get("/api/scan/records", params={"source": kind}))["records"]}


def age(kind, source_id, days):
    """Запись как будто проверена источником days дней назад (прошло время)."""
    with SessionLocal() as session:
        record = session.query(ScanRecord).filter_by(source=kind, source_id=source_id).one()
        record.checked_at = datetime.now(timezone.utc) - timedelta(days=days)
        session.commit()


def antivirus_column(client):
    ws = load_workbook(BytesIO(client.get("/api/export/computers.xlsx").content))["Компьютеры"]
    headers = [cell.value for cell in ws[1]]
    return {row[headers.index("HOSTNAME")].value: row[headers.index("Антивирус")].value for row in ws.iter_rows(min_row=2)}


# ---------- Записи не стираются ----------


def test_records_kept(admin, editor, reader, room, glpi_url):
    loc = room["room"]
    first = add_pc(editor, loc, "pc-1", mac=MAC[1])
    second = add_pc(editor, loc, "pc-2", mac=MAC[2])
    Glpi.computers = {1: pc("pc-1", ports=port(1)), 2: pc("pc-2", ports=port(2, "10.0.9.2"))}
    Glpi.antivirus = [{"id": 1, "computers_id": 2, "name": "KES", "is_active": 1, "is_uptodate": 1}]
    setup_source(admin, glpi_url, "glpi")
    run = collect(admin)
    assert run["status"] == "ok" and run["stats"]["kept"] == 0 and run["stats"]["key"] == 2
    before = diffs_of(editor)
    assert {(d["computer_id"], d["field"]) for d in before["items"]} >= {(first, "ram"), (second, "ram")}
    assert before["antivirus"][str(second)][0]["name"] == "KES" and antivirus_column(reader)["pc-2"] == "KES"

    # Сбой: GLPI отдал 0 ПК – записи и предложения прежние, сбор не с ошибкой
    Glpi.computers = {}
    run = collect(admin)
    assert run["status"] == "ok" and run["stats"]["total"] == 0 and run["stats"]["kept"] == 2
    found = records(editor)
    assert set(found) == {1, 2} and all(r["active"] for r in found.values())
    assert diffs_of(editor)["items"] == before["items"]

    # ПК 2 удалили в GLPI, проверялся давно – запись остаётся давней: сопоставлена, но ничего
    # не предлагает (ни значения, ни антивирус, ни адрес для Vacuum)
    with SessionLocal() as session:
        session.add(ScanJabberUser(login="ivanov", groups=[], online=False, registered=True, last_ip="10.0.9.2"))
        session.commit()
    hosts = ok(editor.get("/api/scan/jabber"))["users"][0]["addresses"][0]["hosts"]
    assert [(h["hostname"], h["source"]) for h in hosts] == [("pc-2", "GLPI №2")]

    Glpi.computers = {1: pc("pc-1", ports=port(1))}
    age("glpi", 2, 10)
    run = collect(admin)
    assert run["stats"]["kept"] == 1 and run["stats"]["key"] == 1 and run["stats"]["total"] == 1
    found = records(editor)
    assert (found[2]["active"], found[2]["state"], found[2]["computer_id"]) == (False, "key", second)
    assert found[1]["active"]
    data, diffs = diffs_data(editor)
    assert (first, "ram") in diffs and not [key for key in diffs if key[0] == second]
    assert str(second) not in data["antivirus"] and antivirus_column(reader)["pc-2"] is None
    assert ok(editor.get("/api/scan/jabber"))["users"][0]["addresses"][0]["hosts"] == []

    # ПК снова проверен – запись обновляется на месте
    Glpi.computers = {2: pc("pc-2", ports=port(2), memory=(16384,))}
    run = collect(admin)
    assert run["stats"]["kept"] == 1
    found = records(editor)
    assert found[2]["active"] and found[2]["values"]["ram"] == "16" and found[1]["active"]


def test_stale_record_yields(admin, editor, room, glpi_url):
    """Агент поставили заново, общих признаков у записей нет: давняя запись – дубль
    действующей, а не конфликт; номер давней в таблице – предлагается новый."""
    loc = room["room"]
    computer = add_pc(editor, loc, "pc-x", mac=MAC[1] + "\n" + MAC[2], glpi_id="5")
    Glpi.computers = {5: pc("pc-x", ports=port(1))}
    setup_source(admin, glpi_url, "glpi")
    collect(admin)
    age("glpi", 5, 10)
    Glpi.computers = {6: pc("pc-x", ports=port(2))}
    collect(admin)

    found = records(editor)
    assert (found[5]["state"], found[5]["dup_of"], found[5]["active"]) == ("dup", 6, False)
    assert (found[6]["state"], found[6]["computer_id"], found[6]["active"]) == ("key", computer, True)
    _, diffs = diffs_data(editor)
    d = diffs[(computer, "glpi_id")]
    assert (d["kind"], d["table"], d["proposed"]) == ("diff", "5", "6")


def test_stale_dups_by_keys(admin, editor, room, glpi_url):
    """Агент на том же железе ставили дважды – живая запись не теряет признаков; давняя с двумя
    MAC не склеивает две живые записи разных ПК."""
    loc = room["room"]
    same = add_pc(editor, loc, "pc-x", mac=MAC[1], serial="SNX12345")
    setup_source(admin, glpi_url, "glpi")
    for number in (1, 2, 3):
        if number > 1:
            age("glpi", number - 1, 10)
        Glpi.computers = {number: pc("pc-x", ports=port(1), serial="SNX12345")}
        run = collect(admin)
    assert run["stats"]["shared"] == 0 and run["stats"]["key"] == 1
    found = records(editor)
    assert [(found[n]["state"], found[n]["dup_of"]) for n in (1, 2)] == [("dup", 3), ("dup", 3)]
    assert (found[3]["state"], found[3]["computer_id"]) == ("key", same)

    first = add_pc(editor, loc, "pc-a", mac=MAC[4])
    second = add_pc(editor, loc, "pc-b", mac=MAC[5])
    Glpi.computers = {7: pc("old-laptop", ports=[(ETH, MAC[4], []), (ETH, MAC[5], [])])}
    collect(admin)
    age("glpi", 7, 30)
    Glpi.computers = {11: pc("pc-a", ports=port(4)), 12: pc("pc-b", ports=port(5))}
    collect(admin)
    found = records(editor)
    assert (found[11]["state"], found[11]["computer_id"]) == ("key", first)
    assert (found[12]["state"], found[12]["computer_id"]) == ("key", second)
    assert found[7]["state"] != "key" and not found[7]["active"]
    _, diffs = diffs_data(editor)
    assert (first, "ram") in diffs and (second, "ram") in diffs


def test_match_stale_rules():
    computers = [
        {"id": 1, "hostname": "pc-1", "ip": "10.0.0.1", "mac": None, "serial": "S1", "glpi_id": "7", "gsit_id": None,
         "archived": False},
        {"id": 2, "hostname": "pc-2", "ip": "10.0.0.2", "mac": None, "serial": None, "glpi_id": None, "gsit_id": None,
         "archived": False},
    ]

    def rec(source_id, name, active, serial=None, ip=None):
        return {"source_id": source_id, "name": name, "keys": {"serial": serial}, "dup_of": None, "active": active,
                "data": {"values": {"ip": ip}}}

    # Давняя по номеру и серийному, действующая по имени – «привязать?» остаётся у действующей
    result, _ = match_all([rec(7, "pc-1", False, "S1"), rec(8, "pc-1", True)], computers, [], "glpi")
    assert result[7]["state"] == "dup" and result[7]["dup_of"] == 8
    assert (result[8]["state"], result[8]["computer_id"]) == ("name", 1)

    # IP давней записи не делает адрес общим; давняя одна – сопоставлена как раньше
    result, _ = match_all([rec(9, "old", False, ip="10.0.0.2"), rec(10, "x", True, ip="10.0.0.2")], computers, [], "glpi")
    assert (result[10]["state"], result[10]["by"], result[10]["computer_id"]) == ("name", ["ip"], 2)
    result, _ = match_all([rec(7, "pc-1", False, "S1")], computers, [], "glpi")
    assert (result[7]["state"], result[7]["computer_id"]) == ("key", 1)

    # Действующая только по IP (чужой ноутбук взял адрес) давнюю запись ПК дублем не делает
    computers[1]["serial"] = None
    result, _ = match_all([rec(7, "pc-1", False), rec(9, "lap-new", True, ip="10.0.0.1")], computers, [], "glpi")
    assert (result[7]["state"], result[7]["computer_id"]) == ("key", 1)
    assert (result[9]["state"], result[9]["by"]) == ("name", ["ip"])


# ---------- Убрать номер GSIT ----------


def collect_both(admin, glpi_url, glpi, gsit):
    setup_source(admin, glpi_url, "glpi")
    setup_source(admin, glpi_url, "gsit")
    Glpi.computers = glpi
    collect(admin, "glpi")
    Glpi.computers = gsit
    collect(admin, "gsit")


def test_gsit_removal(admin, editor, reader, room, glpi_url):
    loc = room["room"]
    moved = add_pc(editor, loc, "pc-1", mac=MAC[1], glpi_id="11", gsit_id="21")
    absent = add_pc(editor, loc, "pc-2", mac=MAC[2], glpi_id="12", gsit_id="99")
    both = add_pc(editor, loc, "pc-3", mac=MAC[3], glpi_id="13", gsit_id="23")
    no_glpi = add_pc(editor, loc, "pc-4", mac=MAC[4], glpi_id="14", gsit_id="24")
    collect_both(
        admin, glpi_url,
        {11: pc("pc-1", ports=port(1)), 12: pc("pc-2", ports=port(2)), 13: pc("pc-3", ports=port(3))},
        {21: pc("pc-1", ports=port(1)), 23: pc("pc-3", ports=port(3)), 24: pc("pc-4", ports=port(4))},
    )
    _, diffs = diffs_data(editor)
    # GSIT на связи – номер не трогаем; записи с номером из таблицы нет – убрать
    assert (moved, "gsit_id") not in diffs and (both, "gsit_id") not in diffs
    d = diffs[(absent, "gsit_id")]
    assert (d["kind"], d["table"], d["proposed"], d["raw"]) == ("diff", "99", "", "–99")
    assert d["sources"][0]["source"] == "gsit" and d["sources"][0]["checked_at"] is None
    assert d["note"] == "нет связи дольше 3 дн., работает GLPI"

    # GSIT замолчал у pc-1 (и у pc-4, где GLPI не на связи) – у pc-1 убрать, у pc-4 нет
    age("gsit", 21, 10)
    age("gsit", 24, 10)
    _, diffs = diffs_data(editor)
    d = diffs[(moved, "gsit_id")]
    assert (d["table"], d["proposed"], d["sources"][0]["source_id"]) == ("21", "", 21)
    assert d["sources"][0]["checked_at"] is not None
    assert (no_glpi, "gsit_id") not in diffs and (both, "gsit_id") not in diffs
    # Других ячеек сканер не очищает: давняя запись GSIT ничего не предлагает
    assert all(d["proposed"] or d["field"] == "gsit_id" for d in diffs.values())

    # Убрать – номер пуст, в Истории – из GSIT; дальше номер не предлагается
    result = ok(editor.post("/api/scan/diffs/accept", json={"items": [
        {"computer_id": moved, "field": "gsit_id", "value": "", "table": "21", "source": "GSIT №21"},
    ]}))
    assert result["accepted"] == 1 and get_row(editor, moved)["gsit_id"] is None
    change = history(editor, "computers", moved)[0]["changes"]["gsit_id"]
    assert (change["old"], change["new"], change["scan"]) == ("21", None, "GSIT №21")
    _, diffs = diffs_data(editor)
    assert (moved, "gsit_id") not in diffs

    # «Оставить как есть» – этот номер этому ПК убрать больше не предлагается
    ok(editor.post("/api/scan/diffs/reject", json={"items": [{"computer_id": absent, "field": "gsit_id", "raw": "–99"}]}))
    data, diffs = diffs_data(editor)
    assert (absent, "gsit_id") not in diffs
    _, rejected = diffs_data(editor, rejected=True)
    assert rejected[(absent, "gsit_id")]["rejected_by"] == "editor"

    # GSIT выключен – о нём ничего не известно, убрать не предлагается
    ok(editor.post("/api/scan/diffs/unreject", json={"items": [{"computer_id": absent, "field": "gsit_id", "raw": "–99"}]}))
    assert (absent, "gsit_id") in diffs_data(editor)[1]
    setup_source(admin, glpi_url, "gsit", enabled=False)
    assert (absent, "gsit_id") not in diffs_data(reader)[1]


def test_removed_gsit_not_offered(admin, editor, room, glpi_url):
    loc = room["room"]
    with_glpi = add_pc(editor, loc, "pc-1", mac=MAC[1], glpi_id="11")
    without = add_pc(editor, loc, "pc-2", mac=MAC[2])
    collect_both(
        admin, glpi_url,
        {11: pc("pc-1", ports=port(1)), 12: pc("pc-2", ports=port(2))},
        {31: pc("pc-1", ports=port(1)), 32: pc("pc-2", ports=port(2))},
    )
    _, diffs = diffs_data(editor)
    assert diffs[(with_glpi, "gsit_id")]["proposed"] == "31" and diffs[(without, "gsit_id")]["proposed"] == "32"

    # Записали и убрали вручную: у ПК с номером GLPI больше не предлагается, у ПК без него – как раньше
    for computer in (with_glpi, without):
        ok(editor.patch(f"/api/computers/{computer}", json={"gsit_id": "31" if computer == with_glpi else "32"}))
        ok(editor.patch(f"/api/computers/{computer}", json={"gsit_id": ""}))
    _, diffs = diffs_data(editor)
    assert (with_glpi, "gsit_id") not in diffs and diffs[(without, "gsit_id")]["proposed"] == "32"

    # Удаление отменили в Истории – номер вернулся; снова убрали – снова не предлагается
    removal = history(editor, "computers", with_glpi)[0]
    ok(editor.post("/api/history/cancel", json={"items": [{"id": removal["id"], "field": "gsit_id"}]}))
    assert get_row(editor, with_glpi)["gsit_id"] == "31"
    ok(editor.patch(f"/api/computers/{with_glpi}", json={"gsit_id": ""}))
    assert (with_glpi, "gsit_id") not in diffs_data(editor)[1]

    # Удаление отменили, а номер потом убрали отменой записи – последняя правка не «убрали»:
    # предлагается снова
    removal = history(editor, "computers", with_glpi)[0]
    ok(editor.post("/api/history/cancel", json={"items": [{"id": removal["id"], "field": "gsit_id"}]}))
    first_set = [h for h in history(editor, "computers", with_glpi) if (h["changes"].get("gsit_id") or {}).get("new") == "31"][-1]
    ok(editor.post("/api/history/cancel", json={"items": [{"id": first_set["id"], "field": "gsit_id"}]}))
    assert get_row(editor, with_glpi)["gsit_id"] is None
    assert diffs_data(editor)[1][(with_glpi, "gsit_id")]["proposed"] == "31"
