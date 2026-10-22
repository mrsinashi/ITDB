"""Этап 69: «Доступ» (логин и пароль) у ПК – как у принтера; права учёток по страницам: роль –
граница, ниже – свои ограничения учётки (нет доступа, просмотр вместо правки); проверка – на
сервере."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

import access_secrets
from auth import PAGES, ROLE_PAGES
from conftest import PASSWORDS, add_computer, add_location, get_row, history, https_client, ok
from db import SessionLocal
from models import User, UserSession
from test_66 import add_move, row_of, types, user_id  # noqa: F401 – types – фикстура
from test_67 import install
from test_printers import add_printer


@pytest.fixture(autouse=True)
def own_perms():
    """Пользователи между тестами не очищаются – свои права учёток после теста снять."""
    yield
    session = SessionLocal()
    try:
        session.query(User).update({User.perms: {}})
        session.commit()
    finally:
        session.close()


def set_perms(admin, login, perms, status=200):
    return ok(admin.patch(f"/api/users/{user_id(login)}", json={"perms": perms}), status)


def second_admin(admin):
    """Ещё один администратор: свои права не меняются – ограничивать его будет admin."""
    ok(admin.post("/api/users", json={"login": "admin2", "role": "admin", "password": "admin2-pass"}))
    client = https_client()
    ok(client.post("/api/auth/login", json={"login": "admin2", "password": "admin2-pass"}))
    return client


@pytest.fixture
def admin2(admin):
    client = second_admin(admin)
    yield client
    session = SessionLocal()
    try:
        uid = session.query(User.id).filter(User.login == "admin2").scalar()
        session.query(UserSession).filter(UserSession.user_id == uid).delete()
        session.query(User).filter(User.id == uid).delete()
        session.commit()
    finally:
        session.close()


# ---------- «Доступ» у ПК ----------


def test_pc_access(admin, editor, reader, room):
    info = ok(reader.get("/api/columns"))
    keys = [column["key"] for column in info["columns"]]
    assert keys.index("web_auth") == keys.index("printers") + 1
    access = next(c for c in info["columns"] if c["key"] == "web_auth")
    # Этап 70: «Доступ» ПК – «Пароль», «Логин» – «Учётная запись», записей несколько (test_70.py)
    assert (access["short"], access["card"], access["kind"]) == ("Пароль", "Пароль", "secret")
    assert "web_auth" in next(g for g in info["card_groups"] if g["title"] == "Сеть")["fields"]
    assert info["history_labels"]["web_password"] == "Пароль"

    pc = add_computer(editor, room["room"], hostname="pc-1")
    assert get_row(reader, pc)["web_auth"] is None
    data = ok(editor.put(f"/api/computers/{pc}/web-auth", json={"entries": [{"login": " Administrator ", "password": "p@ss 1"}]}))
    assert data["updated"]["web_auth"] == "есть"
    assert get_row(reader, pc)["web_auth"] == "есть"
    # В строке и в Истории – без значений; не отменяется
    record = history(editor, "computers", pc)[0]
    assert record["changes"] == {"web_login": {"old": None, "new": "задана новая"}, "web_password": {"old": None, "new": "задан новый"}}
    response = editor.post("/api/history/cancel", json={"items": [{"id": record["id"]}]})
    assert response.status_code == 400
    # Обычной правкой поля не задать; читателю – нельзя
    version = get_row(editor, pc)["version"]
    assert editor.patch(f"/api/computers/{pc}", json={"web_auth": "x", "_version": version}).status_code == 400
    assert reader.put(f"/api/computers/{pc}/web-auth", json={"entries": [{"password": "x"}]}).status_code == 403

    # Показать – по паролю администратора, как у принтера
    url = f"/api/computers/{pc}/web-auth/show"
    assert ok(admin.post(url, json={"password": PASSWORDS["admin"]})) == {"entries": [{"login": "Administrator", "password": "p@ss 1"}]}
    assert editor.post(url, json={"password": PASSWORDS["editor"]}).status_code == 403
    assert ok(reader.post(url, json={"login": "admin", "password": PASSWORDS["admin"]}))["entries"][0]["login"] == "Administrator"
    # Попытки – общие с принтерами: пять неверных у ПК – и принтер не покажет
    printer = add_printer(editor, room["room"])
    for _ in range(5):
        reader.post(url, json={"login": "admin", "password": "wrong"})
    assert reader.post(f"/api/printers/{printer}/web-auth/show", json={"login": "admin", "password": PASSWORDS["admin"]}).status_code == 429
    access_secrets.show_attempts.clear()

    # В выгрузку не идёт
    ws = load_workbook(BytesIO(admin.get("/api/export/computers.xlsx").content))["Компьютеры"]
    assert not {"Доступ", "Пароль"} & {cell.value for cell in ws[1]}

    # Пустая запись – удалить
    ok(editor.put(f"/api/computers/{pc}/web-auth", json={"entries": [{"login": "", "password": ""}]}))
    assert get_row(reader, pc)["web_auth"] is None
    assert set(history(editor, "computers", pc)[0]["changes"]) == {"web_login", "web_password"}


# ---------- Права по страницам ----------


def test_role_levels(admin, editor, reader):
    for client, role in ((admin, "admin"), (editor, "editor"), (reader, "reader")):
        assert ok(client.get("/api/auth/me"))["perms"] == ROLE_PAGES[role]
    # Как было до этапа: у читателя нет правки, Vacuum и «Сканера»; у редактора – «Названий» (с этапа
    # 72 – «Справочники → Значения»), «Расписания», «Подключений», «Пользователей»
    assert set(ROLE_PAGES["reader"].values()) == {"view", "none"}
    assert [page for page in PAGES if ROLE_PAGES["editor"][page] == "none"] == [
        "scan/schedule", "scan/settings", "choices/values", "users"]

    pages = ok(admin.get("/api/users/pages"))
    assert [p["key"] for p in pages] == list(PAGES)
    net = next(p for p in pages if p["key"] == "scan/net")
    # Этап 70: у «Сети» нечего править – только запуск (сбор); роль – граница, base – без своих настроек
    assert (net["group"], net["label"], net["levels"]) == ("Сканер", "Сеть", ["none", "view", "run"])
    assert net["base"] == {"admin": "run", "editor": "view", "reader": "none"}
    assert net["max"] == {"admin": "run", "editor": "run", "reader": "run"}
    assert editor.get("/api/users/pages").status_code == 403


def test_restrict_editor(admin, editor, reader, room):
    # Правка «Компьютеров» снята: смотреть можно, менять – нет; в Историю – словами
    saved = set_perms(admin, "editor", {"table": "view"})
    assert saved["perms"]["table"] == "view" and saved["perms"]["printers"] == "edit"
    record = history(admin, "users", user_id("editor"))[0]
    assert record["changes"] == {"perms": {"old": "Компьютеры: правка", "new": "Компьютеры: просмотр"}}
    assert admin.post("/api/history/cancel", json={"items": [{"id": record["id"]}]}).status_code == 400
    assert ok(editor.get("/api/auth/me"))["perms"]["table"] == "view"

    response = editor.post("/api/computers", json={"location_id": room["room"]})
    assert response.status_code == 403
    assert response.json()["detail"] == "Недостаточно прав: «Компьютеры» – только просмотр."
    ok(editor.get("/api/computers"))
    assert editor.get("/api/export/computers.xlsx").status_code == 200
    add_printer(editor, room["room"])

    # Нет доступа: страница и её выгрузка закрыты; общие данные (список ПК) – нет
    set_perms(admin, "editor", {"table": "none", "history": "none"})
    response = editor.get("/api/export/computers.xlsx")
    assert response.status_code == 403 and response.json()["detail"] == "Нет доступа: «Компьютеры»."
    assert editor.get("/api/history").status_code == 403
    ok(editor.get("/api/computers"))

    # «Как у роли» – ограничение снимается; больше границы роли не дать (этап 70: читателю – правку)
    assert set_perms(admin, "editor", {"table": "edit", "history": "edit"})["perms"] == ROLE_PAGES["editor"]
    add_computer(editor, room["room"])
    assert "не даёт" in set_perms(admin, "reader", {"table": "edit"}, 400)["detail"]
    assert reader.post("/api/computers", json={"location_id": room["room"]}).status_code == 403

    # Неизвестное – нельзя; свои права – тоже
    assert set_perms(admin, "editor", {"nope": "view"}, 400)
    assert set_perms(admin, "editor", {"table": "read"}, 400)
    assert "Свои права" in set_perms(admin, "admin", {"table": "view"}, 400)["detail"]


def test_admin_limits(admin, admin2, room):
    set_perms(admin, "admin2", {"tree": "view", "scan/check": "view", "scan/schedule": "view", "scan/settings": "view"})
    assert admin2.post("/api/locations", json={"kind": "building", "name": "Новый"}).status_code == 403
    # Сбор – с запуском (до этапа 70 – с правкой) там, где его запускают, в «Расписании» или «Подключениях»
    response = admin2.post("/api/scan/sources/glpi/collect")
    assert response.status_code == 403 and "GLPI / GSIT" in response.json()["detail"]
    assert admin2.patch("/api/scan/sources/glpi", json={"url": "http://x"}).status_code == 403
    # «Пользователи» не ограничены: может вернуть права другим
    ok(admin2.get("/api/users"))


def test_history_undo_by_page(admin, editor, room):
    pc = add_computer(editor, room["room"], hostname="pc-1")
    ok(editor.patch(f"/api/computers/{pc}", json={"ram": "8", "_version": get_row(editor, pc)["version"]}))
    record = history(editor, "computers", pc)[0]
    # Правку «Компьютеров» сняли – ни Ctrl+Z, ни отмены в Истории
    set_perms(admin, "editor", {"table": "view"})
    response = editor.post("/api/history/undo-last", json={})
    assert response.status_code == 403 and "Компьютеры" in response.json()["detail"]
    assert editor.post("/api/history/cancel", json={"items": [{"id": record["id"], "field": "ram"}]}).status_code == 403
    assert editor.post("/api/history/revert", json={"id": record["id"], "field": "ram", "initial": True}).status_code == 403
    # «История» – только просмотр, а «Компьютеры» – правка: отмена своей правки – можно (запрет
    # «Истории» прячет кнопки страницы, данные защищает право на сам объект)
    set_perms(admin, "editor", {"table": "edit", "history": "view"})
    ok(editor.post("/api/history/cancel", json={"items": [{"id": record["id"], "field": "ram"}]}))
    assert get_row(editor, pc)["ram"] is None


def test_scanner_and_vacuum_pages(admin, editor):
    ok(editor.get("/api/scan/records", params={"source": "glpi"}))
    ok(editor.get("/api/scan/hosts"))
    ok(editor.get("/api/scan/jabber"))
    set_perms(admin, "editor", {"scan/check": "view", "scan/net": "none", "vacuum": "none"})
    ok(editor.get("/api/scan/records", params={"source": "glpi"}))
    assert editor.get("/api/scan/hosts").status_code == 403
    assert editor.get("/api/scan/jabber").status_code == 403
    # Названия – из «GLPI / GSIT» и из Таблицы («Своё», «для всех ПК»); решения по записям – по-
    # прежнему только админу
    body = {"field": "os", "source": "Windows 10 Pro", "table": "Win 10", "kind": "same"}
    ok(editor.post("/api/scan/names", json=body))
    set_perms(admin, "editor", {"table": "view"})
    assert editor.post("/api/scan/names", json=dict(body, kind="keep")).status_code == 403
    set_perms(admin, "editor", {"scan/check": "edit"})
    ok(editor.post("/api/scan/names", json=dict(body, kind="keep")))
    assert editor.post("/api/scan/records/glpi/1/link", json={"computer_id": 1}).status_code == 403
    # Отмена своего названия – и с одной правкой «Компьютеров»
    set_perms(admin, "editor", {"scan/check": "view", "table": "edit"})
    ok(editor.post("/api/history/undo-last", json={}))


def test_room_move_needs_table(admin, editor, room):
    """Перенос кабинета из ячейки: «только ПК» меняет место ПК – нужна правка «Компьютеров»;
    весь кабинет – только «Структура»."""
    pc = add_computer(editor, room["room"], hostname="pc-1")
    other = add_location(admin, "department", "Хирургия", room["building"])
    set_perms(admin, "editor", {"table": "view"})
    url = f"/api/locations/{room['room']}/move"
    response = editor.post(url, json={"parent_id": other, "computer_id": pc})
    assert response.status_code == 403 and "Компьютеры" in response.json()["detail"]
    assert get_row(admin, pc)["department"] == "Терапия"
    ok(editor.post(url, json={"parent_id": other, "whole": True}))
    assert get_row(admin, pc)["department"] == "Хирургия"


def test_undo_from_other_page(admin, editor, room, models):
    """Картриджи модели правят и ячейкой на «Принтерах»: своё Ctrl+Z – и без правки моделей."""
    printer = add_printer(editor, room["room"], model_id=models["printer"]["id"])

    def set_cartridges(value):
        version = next(r for r in ok(editor.get("/api/printers"))["rows"] if r["id"] == printer)["version"]
        ok(editor.patch(f"/api/printers/{printer}", json={"cartridges": value, "_version": version}))

    # Картриджи – уже в справочнике: новый создаётся без отмены, Ctrl+Z бы на нём остановился
    ok(admin.post("/api/cartridges", json={"name": "CF259A"}))
    ok(admin.post("/api/cartridges", json={"name": "CF259X"}))
    set_perms(admin, "editor", {"choices/models": "view"})
    set_cartridges("CF259A")
    assert history(admin, "printer_models", models["printer"]["id"])[0]["changes"]["cartridges"]["new"] == "CF259A"
    ok(editor.post("/api/history/undo-last", json={}))
    # Без правки и «Принтеров» – уже нельзя
    set_cartridges("CF259X")
    set_perms(admin, "editor", {"printers": "view", "choices/cartridges": "view"})
    assert editor.post("/api/history/undo-last", json={}).status_code == 403


def test_role_change_drops_stale_limits(admin, editor):
    """Смена роли: ограничения не ниже новой роли не хранятся – при возврате роли не оживают."""
    set_perms(admin, "editor", {"table": "view", "history": "none"})
    uid = user_id("editor")
    ok(admin.patch(f"/api/users/{uid}", json={"role": "reader"}))
    ok(admin.patch(f"/api/users/{uid}", json={"role": "editor"}))
    perms = ok(editor.get("/api/auth/me"))["perms"]
    assert (perms["table"], perms["history"]) == ("edit", "none")


def test_journal_pages(admin, editor, room, types):
    monitor = add_move(editor, types["Монитор"], model="Asus")
    set_perms(admin, "editor", {"journals/moves": "view"})
    ok(editor.get("/api/moves"))
    assert editor.post("/api/moves", json={"type_id": types["Монитор"]}).status_code == 403
    set_perms(admin, "editor", {"journals/moves": "none"})
    assert editor.get("/api/moves").status_code == 403

    # «Установить» с заменой ПК – ПК уходит в архив: нужна правка «Компьютеров»
    set_perms(admin, "editor", {"journals/moves": "edit", "table": "view"})
    pc = add_computer(admin, room["room"], hostname="old")
    computer = add_move(editor, types["Компьютер"], model="DEXP")
    response = install(editor, computer, 403, dest=f"l:{room['room']}", replace_pc=pc)
    assert "Компьютеры" in response["detail"]
    assert get_row(admin, pc)["archived"] is False
    install(editor, monitor, dest=f"l:{room['room']}")
    assert row_of(editor, monitor)["move_status"] == "Установлен"


def test_names_and_models_pages(admin, editor, room):
    pc = add_computer(admin, room["room"], hostname="own-name")
    set_perms(admin, "editor", {"table": "view", "printers": "view", "journals/moves": "view"})
    # «Своё имя» – и из «Справочников → Имена ПК»; переименовать ПК – правка «Компьютеров»
    ok(editor.post("/api/naming/keep", json={"ids": [pc], "keep": True, "kind": "pc"}))
    assert editor.post("/api/naming/rename", json={"items": [{"id": pc, "hostname": "x"}], "kind": "pc"}).status_code == 403
    # Модель – из «Справочников», «Принтеров» или журнала
    ok(editor.post("/api/printer-models", json={"kind": "printer", "maker": "HP", "model": "M404"}))
    set_perms(admin, "editor", {"choices/models": "view", "choices/names": "view"})
    assert editor.post("/api/printer-models", json={"kind": "printer", "maker": "HP", "model": "M405"}).status_code == 403
    assert editor.post("/api/naming/keep", json={"ids": [pc], "keep": False, "kind": "pc"}).status_code == 403
