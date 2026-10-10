"""Этап 44: принтеры и МФУ – справочник моделей, № в кабинете (у МФУ и принтеров свой
счёт), подключение к ПК (USB / сеть), логин и пароль Web по паролю администратора,
архив, массовая правка, выгрузка, отмена из Истории."""
from io import BytesIO

import pytest
from cryptography.fernet import Fernet
from openpyxl import load_workbook

import api_printers
from conftest import PASSWORDS, add_computer, add_location, get_row, ok

KEY = Fernet.generate_key().decode()


@pytest.fixture(autouse=True)
def secret_key(monkeypatch):
    monkeypatch.setenv("ITDB_SECRET_KEY", KEY)
    api_printers.show_attempts.clear()


def add_model(client, model, kind="printer", maker=None, color=False, duplex=True):
    return ok(client.post("/api/printer-models", json={"kind": kind, "maker": maker, "model": model, "color": color, "duplex": duplex}))


def add_printer(client, location_id, **fields):
    payload = {"location_id": location_id}
    payload.update(fields)
    return ok(client.post("/api/printers", json=payload))["id"]


def printer_rows(client, archived="no"):
    return ok(client.get("/api/printers", params={"archived": archived}))["rows"]


def printer(client, printer_id):
    return next(row for row in printer_rows(client, "all") if row["id"] == printer_id)


def patch(client, printer_id, status=200, **fields):
    fields["_version"] = printer(client, printer_id)["version"]
    return ok(client.patch(f"/api/printers/{printer_id}", json=fields), status)


def history(client, entity, entity_id):
    return ok(client.get("/api/history", params={"entity": entity, "entity_id": entity_id}))["items"]


@pytest.fixture
def models(editor):
    return {
        "mfu": add_model(editor, "ECOSYS M2040dn", "mfu", "Kyocera"),
        "printer": add_model(editor, "LaserJet M404", "printer", "HP", color=True, duplex=False),
    }


def test_models_directory(editor, reader, models):
    assert reader.post("/api/printer-models", json={"model": "X"}).status_code == 403
    items = ok(reader.get("/api/printer-models"))["items"]
    # Сначала МФУ, потом принтеры
    assert [(m["title"], m["kind"], m["color"], m["duplex"], m["printers"]) for m in items] == [
        ("Kyocera ECOSYS M2040dn", "mfu", False, True, 0),
        ("HP LaserJet M404", "printer", True, False, 0),
    ]
    # Тот же производитель и модель – нельзя; тип – только МФУ / принтер; без названия – нельзя
    assert editor.post("/api/printer-models", json={"maker": "kyocera", "model": " ecosys  m2040dn"}).status_code == 400
    assert editor.post("/api/printer-models", json={"kind": "fax", "model": "X"}).status_code == 400
    assert editor.post("/api/printer-models", json={"model": " "}).status_code == 400

    # Правка модели – в Истории модели; удалить можно, пока нет принтеров
    ok(editor.patch(f"/api/printer-models/{models['printer']['id']}", json={"duplex": True}))
    records = history(reader, "printer_models", models["printer"]["id"])
    assert records[0]["changes"] == {"duplex": {"old": False, "new": True}}
    assert records[0]["title"] == "HP LaserJet M404"
    spare = add_model(editor, "Spare")
    ok(editor.delete(f"/api/printer-models/{spare['id']}"))
    assert len(ok(reader.get("/api/printer-models"))["items"]) == 2


def test_create_and_numbers(editor, reader, room, models):
    assert reader.post("/api/printers", json={"location_id": room["room"]}).status_code == 403
    assert editor.post("/api/printers", json={}).status_code == 400

    # Без № – следующий свободный; у МФУ и принтеров – свой счёт
    p1 = add_printer(editor, room["room"], name="ter-ord-printer-1", model="HP LaserJet M404", ip="10.0.5.20")
    m1 = add_printer(editor, room["room"], name="ter-ord-mfu-1", model="ecosys m2040dn")
    p2 = add_printer(editor, room["room"], model_id=models["printer"]["id"])
    rows = printer_rows(reader)
    # Сначала МФУ, потом принтеры по №
    assert [(r["id"], r["kind"], r["number"]) for r in rows] == [(m1, "МФУ", 1), (p1, "Принтер", 1), (p2, "Принтер", 2)]
    row = rows[1]
    assert (row["building"], row["department"], row["room_code"], row["room_name"]) == ("ул. Ленина, 1", "Терапия", "201", "Ординаторская")
    assert (row["maker"], row["model"], row["color"], row["duplex"], row["model_title"]) == ("HP", "LaserJet M404", "цветная", "нет", "HP LaserJet M404")

    # Занятый № – следующие подряд сдвигаются (как № места у ПК)
    p3 = add_printer(editor, room["room"], number=1, model="HP LaserJet M404")
    numbers = {r["id"]: r["number"] for r in printer_rows(reader)}
    assert (numbers[p3], numbers[p1], numbers[p2], numbers[m1]) == (1, 2, 3, 1)
    assert history(reader, "printers", p1)[0]["changes"] == {"number": {"old": 1, "new": 2}}

    # Неизвестная модель – подсказка, где её добавить
    response = editor.post("/api/printers", json={"location_id": room["room"], "model": "Canon X"})
    assert response.status_code == 400 and "Справочники" in response.json()["detail"]


def test_patch_fields(editor, reader, room, models):
    pid = add_printer(editor, room["room"], model="HP LaserJet M404")
    assert reader.patch(f"/api/printers/{pid}", json={"name": "x"}).status_code == 403

    data = patch(editor, pid, name=" prn-1 ", ip="10.0.5.20, 10.0.5.21", web="есть", web_url="10.0.5.20:8080", inv_no=" 123 ", note="a\n\nb")
    row = data["updated"]
    assert (row["name"], row["ip"], row["web"], row["web_url"], row["inv_no"], row["note"]) == (
        "prn-1", "10.0.5.20\n10.0.5.21", "есть", "http://10.0.5.20:8080", "123", "a\nb")
    assert patch(editor, pid, 400, web="может быть")
    # Значения модели у принтера не правятся – меняется модель
    assert editor.patch(f"/api/printers/{pid}", json={"maker": "Canon"}).status_code == 400
    # Чужая версия – 409
    assert editor.patch(f"/api/printers/{pid}", json={"name": "y", "_version": 1}).status_code == 409

    # Смена типа (другая модель): № в новом счёте занят – следующий свободный
    mfu = add_printer(editor, room["room"], model="Kyocera ECOSYS M2040dn")
    assert printer(reader, mfu)["number"] == 1 and printer(reader, pid)["number"] == 1
    data = patch(editor, pid, model="Kyocera ECOSYS M2040dn")
    assert data["changes"]["number"] == {"old": 1, "new": 2}
    assert data["changes"]["model"]["new"] == "Kyocera ECOSYS M2040dn"
    assert data["updated"]["kind"] == "МФУ"

    # Переезд в другой кабинет – № остаётся, если свободен
    other = add_location(editor, "room", "Процедурная", room["department"], code="202")
    data = patch(editor, pid, location_id=other)
    assert data["updated"]["room_code"] == "202" and data["updated"]["number"] == 2


def test_links_with_computers(editor, reader, room, models):
    pc1 = add_computer(editor, room["room"], hostname="ter-ord-1")
    pc2 = add_computer(editor, room["room"], hostname="ter-ord-2")
    pid = add_printer(editor, room["room"], name="ter-ord-printer-1", model="HP LaserJet M404")

    # У принтера – «Компьютеры»: имя или «имя [usb]», без учёта регистра
    data = patch(editor, pid, computers="TER-ORD-2 [USB]\nter-ord-1")
    assert data["updated"]["computers"] == "ter-ord-1\nter-ord-2 [usb]"
    assert data["updated"]["computer_refs"] == [
        {"id": pc1, "title": "ter-ord-1", "usb": False}, {"id": pc2, "title": "ter-ord-2", "usb": True}]
    # У ПК – «Принтеры»
    assert get_row(reader, pc2)["printers"] == "ter-ord-printer-1 [usb]"
    assert get_row(reader, pc2)["printer_refs"] == [{"id": pid, "title": "ter-ord-printer-1", "usb": True}]
    # В Истории – текстом и номерами записей (отмена – по номерам)
    assert history(reader, "printers", pid)[0]["changes"]["computers"] == {
        "old": None, "new": "ter-ord-1\nter-ord-2 [usb]", "old_ids": [], "new_ids": [[pc1, False], [pc2, True]]}

    # Правка у ПК: по сети вместо USB; в Истории ПК
    version = get_row(reader, pc2)["version"]
    data = ok(editor.patch(f"/api/computers/{pc2}", json={"printers": "ter-ord-printer-1", "_version": version}))
    assert data["updated"]["printers"] == "ter-ord-printer-1"
    assert printer(reader, pid)["computers"] == "ter-ord-1\nter-ord-2"
    change = history(reader, "computers", pc2)[0]["changes"]["printers"]
    assert (change["old"], change["new"], change["new_ids"]) == ("ter-ord-printer-1 [usb]", "ter-ord-printer-1", [[pid, False]])

    # Неизвестное имя, повтор имени – ошибка; «#N» – по номеру записи
    assert patch(editor, pid, 400, computers="нет-такого")
    add_computer(editor, room["room"], hostname="ter-ord-1")
    assert "нескольких" in patch(editor, pid, 400, computers="ter-ord-1")["detail"]
    data = patch(editor, pid, computers=f"#{pc1} [usb]")
    assert data["updated"]["computers"] == "ter-ord-1 [usb]"

    # Пусто – отключить от всех
    assert patch(editor, pid, computers="")["updated"]["computers"] is None
    assert get_row(reader, pc1)["printers"] is None


def test_web_auth(editor, reader, admin, room, models):
    pid = add_printer(editor, room["room"], model="HP LaserJet M404")
    assert reader.put(f"/api/printers/{pid}/web-auth", json={"login": "a", "password": "b"}).status_code == 403
    data = ok(editor.put(f"/api/printers/{pid}/web-auth", json={"login": " admin ", "password": "p@ss 1"}))
    assert data["updated"]["web_auth"] == "есть"
    # Наружу – только «есть»; в базе – зашифровано; в Истории – без значений
    assert "admin" not in str(printer(reader, pid))
    assert history(reader, "printers", pid)[0]["changes"] == {
        "web_login": {"old": None, "new": "задан новый"}, "web_password": {"old": None, "new": "задан новый"}}

    # Показать – только по логину и паролю администратора (кто бы ни смотрел)
    url = f"/api/printers/{pid}/web-auth/show"
    assert reader.post(url, json={"login": "editor", "password": PASSWORDS["editor"]}).status_code == 403
    assert reader.post(url, json={"login": "admin", "password": "wrong"}).status_code == 403
    shown = ok(reader.post(url, json={"login": "ADMIN", "password": PASSWORDS["admin"]}))
    assert shown == {"login": "admin", "password": "p@ss 1"}

    # Пять неверных попыток подряд – дальше отказ
    for _ in range(5):
        reader.post(url, json={"login": "admin", "password": "wrong"})
    assert reader.post(url, json={"login": "admin", "password": PASSWORDS["admin"]}).status_code == 429
    api_printers.show_attempts.clear()

    # Удалить пароль; логин остаётся
    ok(editor.put(f"/api/printers/{pid}/web-auth", json={"password": ""}))
    assert ok(reader.post(url, json={"login": "admin", "password": PASSWORDS["admin"]})) == {"login": "admin", "password": None}
    assert history(reader, "printers", pid)[0]["changes"] == {"web_password": {"old": None, "new": "удалён"}}


def test_archive_bulk_export(editor, reader, room, models):
    a = add_printer(editor, room["room"], name="prn-a", model="HP LaserJet M404")
    b = add_printer(editor, room["room"], name="prn-b", model="HP LaserJet M404")

    ok(editor.post("/api/printers/bulk-update", json={"ids": [a, b], "field": "web", "value": "нет"}))
    assert {printer(reader, a)["web"], printer(reader, b)["web"]} == {"нет"}
    assert editor.post("/api/printers/bulk-update", json={"ids": [a], "field": "name", "value": "x"}).status_code == 400

    # Модель с принтерами (и в архиве) не удаляется
    assert editor.delete(f"/api/printer-models/{models['printer']['id']}").status_code == 400

    ok(editor.post("/api/printers/archive", json={"ids": [b], "archived": True}))
    assert [r["id"] for r in printer_rows(reader)] == [a]
    assert [r["id"] for r in printer_rows(reader, "yes")] == [b]

    # Узел с рабочими принтерами в архив не убрать
    response = editor.post(f"/api/locations/{room['room']}/archive")
    assert response.status_code == 400 and "принтеры" in response.json()["detail"]

    response = reader.get("/api/export/printers.xlsx", params={"archive": "true"})
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["Принтеры", "Архив"]
    header = [cell.value for cell in wb["Принтеры"][1]]
    assert "Вход" not in header and header[:2] == ["Адрес", "Отделение"]
    assert wb["Принтеры"].cell(row=2, column=header.index("Имя") + 1).value == "prn-a"
    assert wb["Архив"].cell(row=2, column=header.index("Имя") + 1).value == "prn-b"


def test_history_undo(editor, reader, room, models):
    pid = add_printer(editor, room["room"], name="prn-1", model="HP LaserJet M404")
    pc = add_computer(editor, room["room"], hostname="pc-1")
    patch(editor, pid, name="prn-2", model="Kyocera ECOSYS M2040dn", computers="pc-1 [usb]")
    record = history(reader, "printers", pid)[0]
    assert set(record["changes"]) == {"name", "model", "computers"}
    assert record["title"] == "prn-2"

    # Ctrl+Z – всё действие разом
    ok(editor.post("/api/history/undo-last"))
    row = printer(reader, pid)
    assert (row["name"], row["model_title"], row["computers"]) == ("prn-1", "HP LaserJet M404", None)
    assert get_row(reader, pc)["printers"] is None

    # Связи отменяются по номерам записей, даже если имя ПК повторяется
    twin = add_computer(editor, room["room"], hostname="pc-1")
    patch(editor, pid, computers=f"#{pc} [usb]")
    patch(editor, pid, computers=f"#{twin}")
    ok(editor.post("/api/history/undo-last"))
    assert printer(reader, pid)["computer_refs"] == [{"id": pc, "title": "pc-1", "usb": True}]

    # Правка модели в справочнике – у всех её принтеров; отмена возвращает
    ok(editor.patch(f"/api/printer-models/{models['printer']['id']}", json={"maker": "Hewlett-Packard"}))
    assert printer(reader, pid)["maker"] == "Hewlett-Packard"
    ok(editor.post("/api/history/undo-last"))
    assert printer(reader, pid)["maker"] == "HP"

    # Логин и пароль Web не отменяются
    ok(editor.put(f"/api/printers/{pid}/web-auth", json={"login": "a"}))
    assert editor.post("/api/history/undo-last").status_code == 400

    # Карточка: строка и история
    card = ok(reader.get(f"/api/printers/{pid}"))
    assert card["row"]["name"] == "prn-1" and card["history"][0]["changes"] == {"web_login": {"old": None, "new": "задан новый"}}
    assert reader.get("/api/printers/99999").status_code == 404


def test_columns_info(reader):
    info = ok(reader.get("/api/columns"))
    keys = [c["key"] for c in info["columns"]]
    # «Принтеры» у ПК – после VNC
    assert keys[keys.index("vnc") + 1] == "printers"
    printer_keys = [c["key"] for c in info["printer_columns"]]
    assert printer_keys[:9] == ["building", "department", "floor", "room_code", "room_name", "number", "name", "computers", "ip"]
    assert {"kind", "maker", "model", "color", "duplex", "web", "web_auth", "inv_no", "serial", "note"} <= set(printer_keys)
